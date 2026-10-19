# -*- coding: utf-8 -*-
"""브랜드/제품 기획서(One-Pager) 시트를 제품개발_마스터.xlsx에 추가.

'남에게 다도기가 뭔지' 설명하는 한 장 — 자사몰 소개·입점 제안·투자 피칭의
공통 뼈대. PPT가 필요하면 이 시트를 슬라이드로 옮기면 된다. 강의 프레임
(화장품기획제조유통.txt: 기획·제품·유통 / 3C+1C / 컨셉·타깃·포지셔닝 /
5가지 제품 역할) 기반으로 다도기 예시를 채워 '고쳐 쓰기'만 하게 설계.

재생성(dev_master.py) 시 GENERATED_SHEETS 밖 시트라 통째 보존됨.
사용: .venv/Scripts/python scripts/brand_plan.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, ValueError):
        pass

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "제품개발_마스터.xlsx"
SHEET = "브랜드 기획서"

DEEP = "3A4A6B"      # 타이틀(딥블루) — OEM(그린)과 구분
SECT = "DCE3F0"      # 섹션 헤더
LABEL = "F1F3F7"     # 항목 라벨
TIPC = "8A8A8A"

WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center")
B = Font(bold=True)
WHITE_B = Font(bold=True, color="FFFFFF", size=13)
TIP = Font(italic=True, color=TIPC, size=9)
thin = Side(style="thin", color="D5DAE2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ("title",t)|("sect",t)|("f",항목,내용예시,팁·프레임)|("note",t)
ROWS = [
    ("title", "다도기 · 브랜드/제품 기획서 (Brand & Product One-Pager)"),
    ("note", "※ 남에게 '다도기가 뭔지' 한 장으로 설명하는 자료. 자사몰 소개·입점 제안·투자 피칭의 공통 뼈대. "
             "파란 글씨는 예시이니 내 값으로 고쳐 쓰세요. PPT가 필요하면 이 섹션들을 슬라이드로 옮기면 됩니다."),

    ("sect", "1. 브랜드 한 줄 (Elevator Pitch)"),
    ("f", "다도기는 ___ 브랜드",
     "지성·예민·지루성 피부를 위한, 자극 유발 성분을 뺀 '무향·세이프 오일' 진정 스킨케어 브랜드",
     "10초 안에 '누구를 위한 무엇'이 들어가게. 가장 많이 쓸 문장."),
    ("f", "태그라인(한 문장)", "“덜어내서 편안한 피부.”", "감성 한 줄. 상세페이지·패키지에 재사용."),

    ("sect", "2. 왜 만드나 (미션 · 스토리)"),
    ("f", "창업자 / 운영", "1인 브랜드, 맞춤형화장품조제관리사 준비, 처방을 직접 설계·검증",
     "'진정성'이 1인 브랜드의 최대 무기. 솔직하게."),
    ("f", "미션", "기름지고 향 강한 진정 크림의 역설을 푼다 — 지성·예민 피부가 안심하고 바르는 제품",
     "'왜 이 브랜드가 세상에 있어야 하나'."),
    ("f", "브랜드 스토리(한 줄)", "내 피부 고민에서 출발 — 맞는 제품이 없어 직접 만들기 시작",
     "출발점의 개인적 계기. 공감 포인트."),

    ("sect", "3. 문제 / 시장의 결핍  (기획 = 데이터로 '결핍' 선점)"),
    ("f", "핵심 결핍",
     "시중 진정·보습 크림 상당수가 말라세지아 먹이(올리브·코코넛·호호바·버터 등 C11~24 지방산 오일)·향료 포함 → 지루성·예민 피부엔 역효과",
     "경쟁이 못 채운 '빈칸'을 한 문장으로."),
    ("f", "데이터 근거(채울 것)", "네이버 데이터랩 '무향/저자극/지성 진정' 검색추이, 화해·올리브영 리뷰의 불만 키워드",
     "강의 Point01: 숫자로 증명. 실제 작성 때 수치·캡처를 넣으면 설득력↑."),

    ("sect", "4. 타겟 고객 (페르소나)"),
    ("f", "핵심 타겟",
     "20~40대, 지성·복합·지루성·예민 피부, 트러블·번들거림·따가움 경험, 전성분 꼼꼼히 보는 소비자",
     "연령·피부·행동을 구체적으로. '모두'는 타겟이 아님."),
    ("f", "고객의 한마디(인용)", "“진정된다는 크림 바르면 오히려 더 번들거리고 좁쌀 올라와요.”",
     "실제 리뷰 톤의 한 줄이 페르소나를 살아있게 함."),

    ("sect", "5. 솔루션 / 핵심 컨셉 (POD · 차별점)"),
    ("f", "3가지 차별점",
     "① 말라세지아-세이프 오일 설계(스쿠알란·MCT·디메치콘) ② 무향 ③ 진정+보습+피지 밸런스",
     "경쟁이 따라오기 어려운 '설계 원칙'으로 묶기."),
    ("f", "한 줄 차별", "“진정 크림인데 유분 걱정이 없다.”", "경쟁 대비 단 한 문장 우위."),

    ("sect", "6. 제품 라인업 · 로드맵  (5가지 제품 역할)"),
    ("f", "히어로(상징)", "시크 진정 밸런스 크림 — 브랜드를 대표하는 완성형", "강의 Point03: 재고 없는 5가지 역할 프레임."),
    ("f", "엔트리(입구)", "토너 / 수분 젤 — 부담 적은 첫 구매", "진입장벽 낮춘 미끼 제품."),
    ("f", "연결(루틴)", "세럼 / 앰플 — 히어로와 함께 쓰는 라인", "객단가·재구매를 올리는 연결고리."),
    ("f", "수익(마진)", "세트 구성 / 리필", "마진·완판을 책임지는 구성."),
    ("f", "실험(확장)", "무수 밤 · 워시오프 클렌징 등 (히어로 검증 후 착수)", "분산 금지 — 히어로 자리잡은 뒤."),

    ("sect", "7. 경쟁 · 포지셔닝  (3C + 1C)"),
    ("f", "Customer(고객)", "위 페르소나 — 지성·예민·성분 중시", "3C+1C: 강의 Part2 프레임."),
    ("f", "Competitor(경쟁)", "시카·장벽 크림 브랜드 다수 — 대부분 향 포함·범용 타겟, '세이프 오일' 소구는 드묾",
     "경쟁의 약점 = 내 기회."),
    ("f", "Company(자사)", "직접 R&D·조제관리사·린 1인 운영·세이프 오일 전문성",
     "작아서 할 수 있는 것(속도·진정성)을 강점으로."),
    ("f", "Channel(채널)", "자사몰 → 올리브영 → 해외 수출(OEM)", "채널별 규격·수수료가 기획에 영향."),
    ("f", "포지셔닝 한 줄", "“지성·예민 전용 × 무향·세이프” = 아직 비어있는 칸 (전문·중가 포지션)",
     "두 축(범용↔전문 / 저가↔고가)에서 내 위치를 한 문장."),

    ("sect", "8. 가격 · 마진  (유통 = 마진 방어)"),
    ("f", "히어로 가격대", "소비자가 2~3만원대 / 목표 원가·마진 ___",
     "심리적 저항선 안에서. Documents\\구매대행 '수출 마진계산기'로 수치화."),
    ("f", "채널 수수료 고려", "올리브영 등 입점 수수료·행사할인 반영한 공급가 설계",
     "마진이 채널에서 깎이는 걸 미리 방어."),

    ("sect", "9. 유통 · 채널 로드맵"),
    ("f", "1단계 — 자사몰", "브랜드 서사·리뷰·재구매 데이터 축적 (가장 통제력 높고 마진 좋음)", "증거(리뷰·판매)를 먼저 쌓기."),
    ("f", "2단계 — 올리브영", "자사몰 성과를 근거로 MD 설득·입점", "강의 Point02: MD 설득엔 데이터가 필요."),
    ("f", "3단계 — 해외 수출", "OEM 생산 → 대상국 규제 대응 후 수출", "관련: 글로벌 수출 방향·마진계산기."),

    ("sect", "10. 왜 나 / 왜 지금 (1인 창업자)"),
    ("f", "강점", "자격(조제관리사)·직접 처방 설계·빠른 의사결정·낮은 고정비·진정성 있는 스토리", "작음을 약점이 아닌 강점으로 번역."),
    ("f", "리스크 & 보완", "경력·자본 한계 → OEM 소량·스톡 용기·자사몰 린스타트로 상쇄", "리스크를 먼저 말하면 신뢰가 는다."),

    ("sect", "11. 마일스톤 (한눈 요약)"),
    ("note", "자격 취득 → 책임판매업 등록 → OEM 샘플·검증 → 히어로 출시(자사몰) → 리뷰·재구매 데이터 축적 → 올리브영/해외 수출"),

    ("sect", "12. 한 줄 요약 (표지 문장)"),
    ("note", "“다도기 — 지성·예민 피부를 위한 무향·세이프 진정 스킨케어. 히어로 ‘시크 진정 밸런스 크림’으로 시작한다.”"),
]


def _title(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text); c.font = WHITE_B
    c.fill = PatternFill("solid", fgColor=DEEP); c.alignment = WRAP_C
    ws.row_dimensions[r].height = 30


def _sect(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text); c.font = B
    c.fill = PatternFill("solid", fgColor=SECT); c.alignment = WRAP_C
    ws.row_dimensions[r].height = 22


def _note(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text)
    c.font = Font(italic=True, size=10, color="3A4A6B")
    c.fill = PatternFill("solid", fgColor="FBFCFE"); c.alignment = WRAP
    c.border = BORDER
    ws.row_dimensions[r].height = 15 * (len(text) // 70 + 1) + 12


def _field(ws, r, label, ex, tip):
    a = ws.cell(r, 1, label); a.font = B
    a.fill = PatternFill("solid", fgColor=LABEL); a.alignment = WRAP
    b = ws.cell(r, 2, ex); b.font = Font(color="1F3B6B"); b.alignment = WRAP
    c = ws.cell(r, 3, tip); c.font = TIP; c.alignment = WRAP
    for col in (1, 2, 3):
        ws.cell(r, col).border = BORDER
    ws.row_dimensions[r].height = 15 * (max(len(ex), len(tip)) // 50 + 1) + 12


def main() -> None:
    if not XLSX.exists():
        sys.exit(f"없음: {XLSX}")
    wb = openpyxl.load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 44
    ws.sheet_view.showGridLines = False

    r = 1
    for row in ROWS:
        kind = row[0]
        if kind == "title":
            _title(ws, r, row[1])
        elif kind == "sect":
            _sect(ws, r, row[1])
        elif kind == "note":
            _note(ws, r, row[1])
        elif kind == "f":
            _field(ws, r, row[1], row[2], row[3])
        r += 1
    ws.freeze_panes = "A4"

    # 'OEM 개발의뢰서' 바로 뒤에 배치
    try:
        anchor = "OEM 개발의뢰서"
        if anchor in wb.sheetnames:
            target = wb.sheetnames.index(anchor) + 1
            wb.move_sheet(SHEET, offset=target - wb.sheetnames.index(SHEET))
    except ValueError:
        pass

    wb.save(XLSX)
    print(f"✅ '{SHEET}' 시트 추가 완료 → {XLSX.name}")
    print("   열: A=항목 / B=내 내용(예시) / C=작성 팁·강의 프레임")


if __name__ == "__main__":
    main()
