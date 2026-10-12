# -*- coding: utf-8 -*-
"""히어로 1배치 제조 플랜 → 엑셀(히어로_배치플랜.xlsx).

sebum-calm-cream v2 처방(formulas/sebum-calm-cream/v2.yaml)에서 배합표(%→그램,
100g 테스트·200g 기준)를 뽑고, 제조순서·QC체크리스트를 붙인다.

⚠ 수치·공정은 설계·연습 기준(예시). 실제 제조 전 원료 TDS·pH·방부력(챌린지)·
사용기한·전성분 검증. 화장품이라 '치료·완화' 표방 금지.

사용: uv run python scripts/hero_batch.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 원료 id → (한글 표기, INCI/비고)
NAME = {
    "water": ("정제수", "Water"),
    "glycerin": ("글리세린", "Glycerin · 습윤"),
    "niacinamide": ("나이아신아마이드", "Niacinamide · 피지·장벽·톤(고시 미백)"),
    "allantoin": ("알란토인", "Allantoin · 진정"),
    "beta-glucan": ("베타글루칸", "Beta-Glucan · 진정·보습"),
    "zinc-pca": ("아연PCA", "Zinc PCA · 피지조절"),
    "pentylene-glycol": ("펜틸렌글라이콜", "Pentylene Glycol · 보습+방부보조"),
    "xanthan-gum": ("잔탄검", "Xanthan Gum · 점증(글리세린 선분산)"),
    "multi-hyaluronic-acid": ("다중 히알루론산", "Hyaluronic Acids · 층별 수분"),
    "montanov-68": ("몬타노브68", "Cetearyl Alcohol & Cetearyl Glucoside · 유화제(세이프)"),
    "cetyl-alcohol": ("세틸알코올", "Cetyl Alcohol · 공유화·바디"),
    "squalane": ("스쿠알란", "Squalane · 경량 에몰리언트(세이프)"),
    "mct": ("MCT오일", "Caprylic/Capric Triglyceride · 경량(세이프)"),
    "dimethicone": ("디메치콘", "Dimethicone · 실키·막"),
    "centella-extract": ("병풀추출물", "Centella Asiatica Extract · 진정·장벽"),
    "panthenol": ("판테놀", "Panthenol(B5) · 보습·진정"),
    "green-tea-extract": ("녹차추출물", "Camellia Sinensis Extract · EGCG 피지·항산화"),
    "phenoxyethanol": ("페녹시에탄올", "Phenoxyethanol · 방부(한도 1%)"),
    "hexanediol": ("1,2-헥산디올", "1,2-Hexanediol · 방부보조·보습"),
}

STEPS = [
    ("0. 소독·준비", "작업대·기구·용기·공병을 70% 에탄올로 소독 후 건조. 니트릴 장갑·마스크. "
     "저울(0.1g)·온도계·pH미터·미니블렌더(호모) 준비."),
    ("1. 계량", "A(수상)·B(유상)·C(냉각) 상을 각각 별도 비커에 계량. "
     "잔탄검은 먼저 글리세린에 개어 두었다가 A상에 넣는다(덩어리 방지)."),
    ("2. A상 가열·용해", "A상을 70~75℃로 가열하며 저어 잔탄검·나이아신·아연PCA·"
     "다중 히알루론산 등 수용성 성분을 완전히 용해(덩어리 없게)."),
    ("3. B상 가열·용해", "B상을 70~75℃로 가열해 몬타노브68·세틸알코올·스쿠알란·MCT·"
     "디메치콘을 투명하게 완전 용해."),
    ("4. 유화", "A·B 온도를 70~75℃로 맞춘 뒤 B를 A에 부으며 미니블렌더로 3~5분 고속 유화. "
     "유백색 크림이 잡히는지 확인."),
    ("5. 냉각 교반", "40℃ 이하로 식히며 저속 교반(찬물 수욕 가능). 점도가 오른다."),
    ("6. C상 투입", "40℃ 이하에서 병풀·판테놀·녹차추출물·페녹시에탄올·1,2-헥산디올을 "
     "넣고 골고루 섞는다(열에 약한 활성·방부 보호)."),
    ("7. pH 측정·조정", "목표 pH 5.0~5.5. 높으면 구연산수(1:9) 소량 적정, 측정 반복."),
    ("8. 충전·라벨", "소독한 50mL 자(jar)에 충전. 제조일·배치번호·보관법 라벨."),
    ("9. 안정성 관찰", "24~48h·1주·4주: 외관(분리·변색)·냄새·pH 재확인. "
     "실온/냉장/고온(45℃) 소량 병행 관찰."),
]

CHECKS = [
    "소독 완료(작업대·기구·용기)",
    "저울 0.1g·온도계·pH미터 준비",
    "잔탄검 글리세린 선분산",
    "A·B 유화 온도 70~75℃(±2)",
    "C상은 40℃ 이하에서 투입",
    "최종 pH 5.0~5.5",
    "외관 균일·분리 없음",
    "패치 테스트(팔 안쪽 48h)",
    "안정성 관찰(실온·냉장·고온 45℃)",
    "★ 방부력(챌린지) 테스트 — 판매 시 필수(외부 의뢰)",
    "기록: 배치번호·일자·실측 그램·이상 메모",
]

NOTE = ("⚠ 설계·연습 기준(예시). 실제 제조 전 원료 TDS·pH·방부력(챌린지)·사용기한·전성분 "
        "검증. 화장품이므로 '치료·완화' 표방 금지(예민·지성용 진정 보습으로 사실 기반 표현).")


def build(formula=Path("formulas/sebum-calm-cream/v2.yaml"),
          out=Path("히어로_배치플랜.xlsx")) -> Path:
    d = yaml.safe_load(formula.read_text(encoding="utf-8"))
    base = d.get("base_batch_g", 200)

    thin = Side(style="thin", color="D8DCE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    hdr = PatternFill("solid", fgColor="E8EEF7")
    phase_fill = PatternFill("solid", fgColor="EFF4FB")

    wb = openpyxl.Workbook()
    # ── 배합표 ──
    ws = wb.active
    ws.title = "배합표"
    ws.merge_cells("A1:F1")
    tt = ws.cell(1, 1, f"{d['product']} (v{d['version']}) — 배합표")
    tt.font = Font(bold=True, size=14, color="33475B")
    ws.merge_cells("A2:F2")
    n = ws.cell(2, 1, NOTE)
    n.font = Font(size=9, color="7A4A2A")
    n.fill = PatternFill("solid", fgColor="FDECE0")
    n.alignment = wrap
    ws.row_dimensions[2].height = 30
    cols = ["상(phase)", "원료", "%", "100g 배치(g)", f"{base}g 기준(g)", "INCI · 역할"]
    widths = [12, 18, 7, 13, 14, 48]
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cc = ws.cell(3, i, c)
        cc.font = Font(bold=True)
        cc.fill = hdr
        cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cc.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 4
    sum_pct = 0.0
    for phase in d["phases"]:
        pname = phase["name"]
        for ing in phase["ingredients"]:
            iid = ing["id"]
            pct = float(ing["percent"])
            sum_pct += pct
            ko, inci = NAME.get(iid, (iid, ""))
            vals = [pname, ko, pct, round(pct, 2), round(pct * base / 100, 1), inci]
            for ci, v in enumerate(vals, 1):
                cc = ws.cell(r, ci, v)
                cc.border = border
                cc.alignment = wrap if ci in (2, 6) else Alignment(
                    horizontal="center", vertical="top")
            ws.cell(r, 1).fill = phase_fill
            r += 1
    # 합계
    ws.cell(r, 1, "합계").font = Font(bold=True)
    ws.cell(r, 3, round(sum_pct, 1)).font = Font(bold=True)
    ws.cell(r, 4, round(sum_pct, 1)).font = Font(bold=True)
    ws.cell(r, 5, round(sum_pct * base / 100, 1)).font = Font(bold=True)
    for ci in range(1, 7):
        ws.cell(r, ci).border = border
        ws.cell(r, ci).fill = hdr
    # 공정 요약(상별)
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "상별 공정 요약").font = Font(bold=True, color="33475B")
    r += 1
    for phase in d["phases"]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
        ws.cell(r, 1, phase["name"]).font = Font(bold=True, size=9)
        ws.cell(r, 1).fill = phase_fill
        ws.cell(r, 1).alignment = wrap
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, phase.get("process", "")).alignment = wrap
        ws.row_dimensions[r].height = 30
        r += 1
    ws.freeze_panes = "A4"

    # ── 제조순서 ──
    ps = wb.create_sheet("제조순서")
    ps.column_dimensions["A"].width = 16
    ps.column_dimensions["B"].width = 96
    ps.cell(1, 1, "제조 순서 (소량 테스트 배치)").font = Font(bold=True, size=13, color="33475B")
    rr = 2
    for title, body in STEPS:
        tc = ps.cell(rr, 1, title)
        tc.font = Font(bold=True, size=10)
        tc.fill = phase_fill
        tc.alignment = wrap
        tc.border = border
        bc = ps.cell(rr, 2, body)
        bc.alignment = wrap
        bc.border = border
        ps.row_dimensions[rr].height = max(30, -(-len(body) // 52) * 15)
        rr += 1

    # ── QC·체크리스트 ──
    qc = wb.create_sheet("QC체크리스트")
    qc.column_dimensions["A"].width = 6
    qc.column_dimensions["B"].width = 70
    qc.cell(1, 1, "☑").font = Font(bold=True)
    qc.cell(1, 2, "제조·품질 체크리스트").font = Font(bold=True, size=13, color="33475B")
    rr = 2
    for item in CHECKS:
        qc.cell(rr, 1, "☐").alignment = Alignment(horizontal="center")
        c = qc.cell(rr, 2, item)
        c.alignment = wrap
        c.border = border
        qc.cell(rr, 1).border = border
        qc.row_dimensions[rr].height = max(20, -(-len(item) // 40) * 16)
        rr += 1
    qc.cell(rr + 1, 2, NOTE).font = Font(size=9, color="7A4A2A")
    qc.cell(rr + 1, 2).alignment = wrap

    wb.save(out)
    print(f"히어로 배치플랜 → {out.name}: 배합표(합 {sum_pct:.1f}%)·제조순서·QC")
    return out


if __name__ == "__main__":
    build()
