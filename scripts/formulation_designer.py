# -*- coding: utf-8 -*-
"""처방 설계 시트: 학습(아키타입 골격·원료) → 창작(직접 설계)의 고리 완성.

인터랙티브 엑셀(수식·드롭다운·자동검증):
 - 제형 선택 → 그 제형의 목표 골격(%)·배합 주의 자동 표시(VLOOKUP)
 - 원료·역할·투입% 입력 → 그램 자동(배치량 기준) + phase별 합계 + 경고
   (합계≠100%, 물+오일인데 유화제 없음, 수상인데 방부 없음)

데이터 출처:
 - 참고_골격 = 카드_아키타입/*.xlsx 치트시트(전형 구성·주의) 추출
 - 원료_고르기 = 원료_사전(천연, 역할별 상위) + ingredient_knowledge(기능성)

사용: uv run python scripts/formulation_designer.py            → 제품개발_마스터.xlsx '처방설계' 시트
      uv run python scripts/formulation_designer.py --standalone → 처방_설계_템플릿.xlsx(구 단독 파일)
"""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

from brandlab.herbnoori_detail import ROLE_COLORS, ROLE_ORDER, glossary_key

try:
    from ingredient_knowledge import PALETTE  # type: ignore
except ImportError:  # 스크립트 디렉터리 기준 import 보조
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    from ingredient_knowledge import PALETTE

THIN = Side(style="thin", color="D0D5DD")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
HDR_FILL = PatternFill("solid", fgColor="E8EEF7")
KEY_FILL = PatternFill("solid", fgColor="FCEFBE")
NOTE_FILL = PatternFill("solid", fgColor="FDECE0")


def _read_archetypes(cards_dir: Path):
    """카드 파일에서 (제형, 골격문자열, 주의) 추출."""
    out = []
    for f in sorted(cards_dir.glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        arche = re.sub(r"^\d+_", "", f.stem)
        ws = openpyxl.load_workbook(f, read_only=True).worksheets[0]
        skel, tips = "", ""
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            v = str(row[0] or "")
            if v.startswith("전형 구성"):
                skel = v.replace("전형 구성(%):", "").strip()
            elif v.startswith("⚠"):
                tips = v.replace("⚠ 이 제형 배합 주의·기본원칙", "").strip().replace("\n", " / ")
        out.append((arche, skel, tips))
    return out


def _palette_rows(glossary: Path, top_per_role: int = 12):
    """원료 고르기용: 천연(사전 역할별 상위) + 기능성 합쳐 역할별로."""
    by_role = defaultdict(list)
    if glossary.exists():
        ws = openpyxl.load_workbook(glossary, read_only=True).worksheets[0]
        # 역할(A) 원료(B) 등장수(C) 대표%(D)
        for row in ws.iter_rows(min_row=2, values_only=True):
            role, name, cnt = row[0], row[1], row[2]
            pct = row[3] if len(row) > 3 else ""
            if role in ROLE_ORDER and name:
                by_role[role].append((str(name), cnt or 0, str(pct or ""), "천연"))
    for role in by_role:
        by_role[role].sort(key=lambda x: -(x[1] or 0))
        by_role[role] = by_role[role][:top_per_role]
    # 기능성 추가
    for role, name, pct, *_ in PALETTE:
        key = role if role in ROLE_ORDER else "기타"
        by_role[key].append((name, 0, pct, "기능성"))
    return by_role


def build(out: Path = Path("처방_설계_템플릿.xlsx")) -> Path:
    arches = _read_archetypes(Path("카드_아키타입"))
    palette = _palette_rows(Path("원료_사전.xlsx"))

    wb = openpyxl.Workbook()

    # ── 시트2 준비: 참고_골격 (VLOOKUP 소스) ─────────────────────────────
    ref = wb.create_sheet("참고_골격")
    for i, t in enumerate(["제형", "목표 골격(전형 %)", "배합 주의"], 1):
        c = ref.cell(1, i, t)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
    ref.column_dimensions["A"].width = 18
    ref.column_dimensions["B"].width = 70
    ref.column_dimensions["C"].width = 70
    for r, (arche, skel, tips) in enumerate(arches, 2):
        ref.cell(r, 1, arche)
        ref.cell(r, 2, skel).alignment = WRAP
        ref.cell(r, 3, tips).alignment = WRAP
        ref.row_dimensions[r].height = 40
    ref.freeze_panes = "A2"
    n_arch = len(arches)

    # ── 시트3: 원료_고르기 ───────────────────────────────────────────────
    pal = wb.create_sheet("원료_고르기")
    for i, t in enumerate(["역할", "원료", "전형/대표%", "구분"], 1):
        c = pal.cell(1, i, t)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
    for col, w in zip("ABCD", (10, 40, 12, 8)):
        pal.column_dimensions[col].width = w
    r = 2
    for role in ROLE_ORDER:
        for name, _cnt, pct, kind in palette.get(role, []):
            pal.cell(r, 1, role).fill = PatternFill(
                "solid", fgColor=ROLE_COLORS.get(role, "F0F0F0"))
            pal.cell(r, 2, name)
            pal.cell(r, 3, pct).alignment = Alignment(horizontal="center")
            pal.cell(r, 4, kind).alignment = Alignment(horizontal="center")
            for ci in range(1, 5):
                pal.cell(r, ci).border = GRID
            r += 1
    pal.freeze_panes = "A2"
    pal.auto_filter.ref = f"A1:D{r - 1}"

    # ── 시트1: 설계 (작업 캔버스) ────────────────────────────────────────
    ws = wb.active
    ws.title = "설계"
    for col, w in zip("ABCDE", (28, 11, 9, 9, 30)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["F"].width = 2
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 10

    ws.merge_cells("A1:E1")
    t = ws.cell(1, 1, "처방 설계 시트  —  학습한 골격으로 직접 설계")
    t.font = Font(bold=True, size=14, color="33475B")

    ws.cell(2, 1, "① 제형 선택").font = Font(bold=True)
    ws.cell(2, 1).fill = KEY_FILL
    ws.cell(2, 2).fill = KEY_FILL
    ws.cell(2, 2).border = GRID
    ws.cell(2, 4, "② 총 배치량(g)").font = Font(bold=True)
    ws.cell(2, 4).fill = KEY_FILL
    ws.cell(2, 5, 100)
    ws.cell(2, 5).fill = KEY_FILL
    ws.cell(2, 5).border = GRID
    ws.cell(2, 5).font = Font(bold=True)

    ws.cell(3, 1, "목표 골격(참고)").font = Font(bold=True, size=9, color="55606E")
    ws.merge_cells("B3:E3")
    ws.cell(3, 2, "=IFERROR(VLOOKUP($B$2,참고_골격!$A:$C,2,FALSE),"
                  "\"← 제형을 선택하세요\")").alignment = WRAP
    ws.row_dimensions[3].height = 32
    ws.cell(4, 1, "배합 주의").font = Font(bold=True, size=9, color="7A4A2A")
    ws.merge_cells("B4:E4")
    ws.cell(4, 2, "=IFERROR(VLOOKUP($B$2,참고_골격!$A:$C,3,FALSE),\"\")").alignment = WRAP
    ws.cell(4, 1).fill = NOTE_FILL
    for cc in range(2, 6):
        ws.cell(4, cc).fill = NOTE_FILL
    ws.row_dimensions[4].height = 40

    # 입력 표 헤더(6행), 입력 7~26행.
    R0, R1 = 7, 26
    for i, h in enumerate(["원료명", "역할", "투입%", "그램", "메모"], 1):
        c = ws.cell(6, i, h)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.border = GRID
        c.alignment = Alignment(horizontal="center")
    for r in range(R0, R1 + 1):
        for ci in range(1, 6):
            ws.cell(r, ci).border = GRID
            ws.cell(r, ci).alignment = WRAP
        # 그램 = 투입% × 총량 /100
        ws.cell(r, 4, f"=IF($C{r}=\"\",\"\",ROUND($C{r}*$E$2/100,2))")
    # 합계 행
    ws.cell(R1 + 1, 1, "합계").font = Font(bold=True)
    ws.cell(R1 + 1, 1).fill = HDR_FILL
    ws.cell(R1 + 1, 3, f"=SUM(C{R0}:C{R1})").font = Font(bold=True)
    ws.cell(R1 + 1, 4, f"=SUM(D{R0}:D{R1})").font = Font(bold=True)
    for ci in range(1, 6):
        ws.cell(R1 + 1, ci).border = GRID

    # phase 합계(우측 G:H).
    ws.cell(6, 7, "phase 합계(%)").font = Font(bold=True)
    ws.merge_cells("G6:H6")
    ws.cell(6, 7).fill = HDR_FILL
    for i, role in enumerate(ROLE_ORDER):
        rr = R0 + i
        gc = ws.cell(rr, 7, role)
        gc.fill = PatternFill("solid", fgColor=ROLE_COLORS.get(role, "F0F0F0"))
        gc.border = GRID
        hc = ws.cell(rr, 8,
                     f'=IF(SUMIF($B${R0}:$B${R1},G{rr},$C${R0}:$C${R1})=0,"",'
                     f'ROUND(SUMIF($B${R0}:$B${R1},G{rr},$C${R0}:$C${R1}),1))')
        hc.border = GRID
        hc.alignment = Alignment(horizontal="center")

    # 자동 검증(우측, phase 아래).
    cr = R0 + len(ROLE_ORDER) + 1
    ws.cell(cr, 7, "자동 점검").font = Font(bold=True)
    ws.merge_cells(start_row=cr, start_column=7, end_row=cr, end_column=8)
    ws.cell(cr, 7).fill = KEY_FILL
    checks = [
        (f'=IF(SUM(C{R0}:C{R1})=0,"→ 원료를 입력하세요",'
         f'IF(ABS(SUM(C{R0}:C{R1})-100)<0.5,"✓ 합계 100%",'
         f'"⚠ 합계 "&ROUND(SUM(C{R0}:C{R1}),1)&"% (100 맞추기)"))'),
        (f'=IF(AND(SUMIF(B{R0}:B{R1},"수상",C{R0}:C{R1})>0,'
         f'SUMIF(B{R0}:B{R1},"유상",C{R0}:C{R1})>0,'
         f'SUMIF(B{R0}:B{R1},"유화",C{R0}:C{R1})=0),'
         f'"⚠ 물+오일인데 유화제 없음 → 분리","")'),
        (f'=IF(AND(SUMIF(B{R0}:B{R1},"수상",C{R0}:C{R1})>0,'
         f'SUMIF(B{R0}:B{R1},"방부",C{R0}:C{R1})=0),'
         f'"⚠ 수상 있는데 방부 없음 → 변질 위험(무방부는 냉장·단기)","")'),
        (f'=IF(SUMIF(B{R0}:B{R1},"향",C{R0}:C{R1})>1,'
         f'"⚠ 향 1% 초과 → 자극/규제 주의","")'),
    ]
    for i, fml in enumerate(checks, 1):
        c = ws.cell(cr + i, 7, fml)
        ws.merge_cells(start_row=cr + i, start_column=7, end_row=cr + i, end_column=8)
        c.alignment = WRAP
        c.font = Font(size=9)
        ws.row_dimensions[cr + i].height = 28

    # 드롭다운: 제형(참고_골격 A열), 역할(인라인).
    dv_arche = DataValidation(
        type="list", formula1=f"=참고_골격!$A$2:$A${n_arch + 1}", allow_blank=True)
    ws.add_data_validation(dv_arche)
    dv_arche.add(ws.cell(2, 2))
    dv_role = DataValidation(
        type="list", formula1='"' + ",".join(ROLE_ORDER) + '"', allow_blank=True)
    ws.add_data_validation(dv_role)
    dv_role.add(f"B{R0}:B{R1}")

    # 사용법 메모(하단).
    ur = cr + len(checks) + 2
    ws.merge_cells(start_row=ur, start_column=1, end_row=ur, end_column=5)
    guide = ("사용법: ① 제형 선택(위 골격·주의 참고) → ② 배치량 입력 → "
             "③ 원료·역할·투입% 입력(원료는 '원료_고르기' 시트 참고). "
             "그램·phase합계·점검은 자동. ⚠ 수치는 설계 연습용 — 실제 제조 전 "
             "식약처 고시·공급사 TDS·방부력 테스트 검증.")
    gc = ws.cell(ur, 1, guide)
    gc.alignment = WRAP
    gc.font = Font(size=9, color="55606E")
    gc.fill = NOTE_FILL
    for cc in range(1, 6):
        ws.cell(ur, cc).fill = NOTE_FILL
    ws.row_dimensions[ur].height = 56

    ws.freeze_panes = "A7"
    wb.save(out)
    print(f"처방 설계 템플릿 → {out.name} (제형 {n_arch}종, 원료고르기 포함)")
    return out


# ── 제품개발_마스터.xlsx에 '처방설계' 시트로 녹이기 ─────────────────────────────
# 단일 워크북 원칙(CLAUDE.md): 설계 캔버스 → 참고 골격 → 학습용 원료(원료_지식_확장·원료_사전)를 한 시트에 위→아래로 쌓는다.
MASTER = Path("제품개발_마스터.xlsx")
SHEET = "처방설계"
NAVY = "FF1F3A5F"
GREY = "FFF2F2F2"
LBLUE = "FFDCE6F1"
GREEN_TAB = "FF548235"  # 🟩 ① 제품 개발 그룹
REDFILL = PatternFill("solid", fgColor="FFFFC7CE")

# 표기 흔들림 정규화(중복 병합·보유 매칭용)
_ALIASES = [("다이올", "디올"), ("알콜", "알코올"), ("쟁탄", "잔탄"), ("히아루론", "히알루론"),
            ("코코-베타인", "코코베타인"), ("에칠", "에틸"), ("피카", "pca"), ("다이메치콘", "디메치콘")]
# 이름에 보유 원료명이 '부분 문자열'로만 들어간 다른 원료(코코베타인≠베타인, 에틸헥실글리세린≠글리세린)
_NOT_OWNED_PARTS = ("코카미도프로필베타인", "코코베타인", "에틸헥실글리세린", "폴리글리세릴")


def _norm(s: str) -> str:
    s = str(s).lower()
    for a, b in _ALIASES:
        s = s.replace(a, b)
    return re.sub(r"[\s/*.,\-&()'’·]", "", s)


# 원료_사전의 자동 분류(이름 키워드)가 '기타'로 떨어뜨렸거나 명백히 틀린 역할 → 학습용 세부 역할.
# 원본 xlsx는 건드리지 않고 시트에 그릴 때만 보정한다(원료_사전 재생성에도 안전).
_EXTRA_ROLES = {  # 학습 역할 12종에 없던 묶음
    "항산화": "FBE9D0", "컨디셔닝": "DDE7F7", "자외선차단": "FFF2B3", "분말·스크럽": "F3E9DC",
    "완제 베이스": "E4E4E4", "장식·건초": "EFE6D8", "생활·공예": "E8E8E8",
    # 기능성 팔레트 역할
    "에몰리언트": "FFF0CC", "감촉·실리콘": "FFF7E0", "분말": "F3E9DC", "가용화": "E3F7E3",
    "킬레이트·pH": "EEF3E2", "자외선": "FFF2B3",
}
_ROLE_FIX = {
    # 기타 → 세부 역할
    **dict.fromkeys(["알파 비사보롤", "수딩 4 컴플렉스", "멘톨", "L-멘톨", "PDRN", "인삼 줄기세포배양액",
                     "인삼줄기 / 세포배양액", "바이오 디톡스", "MSM (식이유황)", "살리실산(0.5%)", "로먼카모마일",
                     "부활초 / 줄기세포 배양액", "재생, 항노화 / 기능성 원료", "징크피리치온", "7가지한방약재",
                     "폴리페놀", "레드와인", "황금", "금은화", "진피", "당귀", "인플라신", "시카엑소좀"], "활성"),
    **dict.fromkeys(["트레할로스"], "보습"),
    **dict.fromkeys(["세틸알코올"], "경화"),
    **dict.fromkeys(["자작나무수액", "편백수"], "수상"),
    **dict.fromkeys(["유기농 달맞이종자유", "피마자유", "마카다미아넛", "골든호호바", "호호바 라이트", "IPM"], "유상"),
    **dict.fromkeys(["그린티 시트러스", "끌로에 우먼", "플라워 바디 겐조", "버가못FCF", "로즈제라늄"], "향"),
    **dict.fromkeys(["천연한방존제"], "방부"),
    **dict.fromkeys(["립컬러 베이스", "립베이스 레드27", "립 컬러베이스 7", "레드 립 컬러", "벚꽃 립 컬러",
                     "피치코랄 립 컬러", "립 컬러", "컬러믹스파운데이션", "페이스파우더 21호", "헤나분말",
                     "물에 갠 검은 색 물감", "칼라믹스"], "색소"),
    **dict.fromkeys(["베이킹 파우더 / (소다)", "덴탈 점도 실리카"], "점증/pH"),
    **dict.fromkeys(["폴리쿼터"], "컨디셔닝"),
    **dict.fromkeys(["시너메이트", "산화아연"], "자외선차단"),
    **dict.fromkeys(["옥수수전분", "오트밀분말", "히말라야 핑크소금", "어성초분말", "카카오분말", "노니분말",
                     "로즈마리 분말", "핑크솔트", "핑크소금", "살구분말", "유노하나", "원두 커피 / 곱게 간 것",
                     "설탕", "제주화산송이", "덴탈 연마 실리카", "덴탈 실리카(연마용)",
                     "히말라아 핑크솔트 / (굵은 소금)"], "분말·스크럽"),
    **dict.fromkeys(["아르간 가득한 / 크림베이스", "허브누리의 촉촉 / 화이트닝 수분크림베이스", "립밤 베이스",
                     "촉촉 화이트닝 / 수분크림베이스", "투명 썬스틱베이스", "미니멈 크림베이스",
                     "에센스 / (4GF 재생 / 에센스)", "허브누리의 촉촉 화이트닝 수분크림베이스", "아르간크림베이스",
                     "블루밍 CC크림", "허브누리 / 순한 샴푸베이스", "시카 판테놀 / 바디&핸드 로션베이스",
                     "허브누리의 아주 순한 샴푸베이스", "워셔블클렌징리퀴드", "워셔블 / 클렌징리퀴드"], "완제 베이스"),
    **dict.fromkeys(["24k 금설", "금설", "24K 금설", "라벤더건초", "캐모마일꽃건초", "장미꽃봉우리", "계피"],
                    "장식·건초"),
    **dict.fromkeys(["친환경 석고", "천연 사쉐 스톤", "친환경 석고 분말", "에코 제습제", "과탄산나트륨",
                     "디퓨저 전용 / 베이스"], "생활·공예"),
    # 다른 역할에 잘못 들어간 것
    "토코페롤": "항산화",
    "세틸에틸헥사노에이트": "유상",
    "로즈플라워 아로마수": "수상", "러블리 레몬 아로마수": "수상",
    "아하바하파하 컴플렉스": "활성",
}
GLOSSARY_ROLE_ORDER = [r for r in ROLE_ORDER if r != "기타"] + [
    "항산화", "컨디셔닝", "자외선차단", "분말·스크럽", "완제 베이스", "장식·건초", "생활·공예", "기타"]


def _role_color(role: str) -> str:
    return ROLE_COLORS.get(role) or _EXTRA_ROLES.get(role, "F0F0F0")


def _read_glossary(path: Path) -> list[tuple]:
    """원료_사전.xlsx → (역할, 원료, 등장수, 대표%, 특징, 대체재, 범위, 주사용제형, 다른 표기). 헤더 이름 기준."""
    if not path.exists():
        return []
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
    want = ["역할", "원료", "등장수", "대표 사용%", "대표 특징", "대체재(모음)", "사용% 범위(최소~중앙~최대)", "주 사용 제형",
            "다른 표기"]
    fix = {glossary_key(k): v for k, v in _ROLE_FIX.items()}  # 표기 변형이 합쳐져도 같은 키로 보정
    out = []
    for row in rows[1:]:
        if not row or not row[idx["원료"]]:
            continue
        rec = [row[idx[k]] if k in idx and idx[k] < len(row) else None for k in want]
        rec[0] = fix.get(glossary_key(rec[1]), rec[0])
        out.append(tuple(rec))
    order = {r: i for i, r in enumerate(GLOSSARY_ROLE_ORDER)}
    out.sort(key=lambda x: (order.get(x[0], len(order)), -int(x[2] or 0) if str(x[2] or "").isdigit() else 0))
    return out


def _read_knowledge(path: Path) -> list[tuple]:
    """원료_지식_확장.xlsx '기능성원료' → (역할, 원료, 전형%, 기능·왜, 천연대안, 주의). 헤더 행 자동 탐지."""
    if not path.exists():
        return [(role, name, pct, "", "", "") for role, name, pct, *_ in PALETTE]
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["기능성원료"] if "기능성원료" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, row in enumerate(rows) if row and row[0] == "역할")
    return [tuple(row[:6]) for row in rows[hdr + 1:] if row and row[1]]


def _master_ingredients(wb) -> list[tuple[str, str, str]]:
    """원료 시트의 (원료, INCI, 세이프)."""
    if "원료" not in wb.sheetnames:
        return []
    ws = wb["원료"]
    out = []
    for r in range(4, ws.max_row + 1):
        ko = ws.cell(r, 1).value
        if ko:
            out.append((str(ko), str(ws.cell(r, 2).value or ""), str(ws.cell(r, 6).value or "")))
    return out


def _match_owned(name: str, owned) -> str:
    """팔레트 원료 → 원료 시트(보유·재고판)에 있는 같은 원료 표시(가장 긴 이름 일치 우선)."""
    n = _norm(name)
    for part in _NOT_OWNED_PARTS:
        n = n.replace(_norm(part), "")
    best, best_len = "", 0
    for ko, inci, safe in owned:
        k = _norm(re.sub(r"\(.*?\)", "", ko))
        keys = {k} | ({_norm(inci)} if inci else set())
        if k.endswith("오일") and len(k) > 4:
            keys.add(k[:-2])  # MCT오일 → mct
        for x in keys:
            if len(x) >= 2 and x in n and len(x) > best_len:
                best, best_len = f"✓ {ko} (세이프 {safe or '?'})", len(x)
    return best


def _x14_validations(path: Path) -> dict[str, list[tuple[str, str]]]:
    """openpyxl이 로드 시 버리는 Excel 확장(x14) 드롭다운을 미리 읽어 둔다 → {시트명: [(formula, sqref)]}."""
    import zipfile
    out: dict[str, list[tuple[str, str]]] = {}
    with zipfile.ZipFile(path) as z:
        wbx = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        rid2file = {m.group(1): m.group(2) for m in re.finditer(
            r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
        rid2file.update({m.group(2): m.group(1) for m in re.finditer(
            r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
        for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
            name, rid = m.group(1), m.group(2)
            target = rid2file.get(rid, "").lstrip("/")
            target = target if target.startswith("xl/") else "xl/" + target
            try:
                xml = z.read(target).decode("utf-8")
            except KeyError:
                continue
            for dv in re.finditer(r"<x14:dataValidation [^>]*type=\"list\".*?</x14:dataValidation>", xml, re.S):
                f = re.search(r"<xm:f>(.*?)</xm:f>", dv.group(0))
                sq = re.search(r"<xm:sqref>(.*?)</xm:sqref>", dv.group(0))
                if f and sq:
                    out.setdefault(name, []).append((f.group(1), sq.group(1)))
    return out


def _insert_row(ws, idx: int) -> None:
    """insert_rows는 병합·행높이를 안 밀어줘서 직접 민다(읽어보기용)."""
    from openpyxl.worksheet.cell_range import CellRange
    merges = [CellRange(str(m)) for m in ws.merged_cells.ranges if m.min_row >= idx]
    for m in merges:
        ws.unmerge_cells(m.coord)
    heights = {i: d.height for i, d in list(ws.row_dimensions.items()) if i >= idx}
    ws.insert_rows(idx)
    for i in sorted(heights, reverse=True):
        ws.row_dimensions[i + 1].height = heights[i]
    ws.row_dimensions[idx].height = None
    for m in merges:
        ws.merge_cells(start_row=m.min_row + 1, start_column=m.min_col,
                       end_row=m.max_row + 1, end_column=m.max_col)


README_LINE = ("처방설계 : 보유 원료로 내 처방을 짜는 캔버스 — 원료만 고르면 역할·상·세이프·현재고·법적한도가 원료 시트에서 따라오고, "
               "그램·역할별 합계·점검(합계100·유화제·방부·세이프X·재고부족·한도초과) 자동. 제형 고르면 목표 골격 비교. 원료 공부는 '원료 도감'.")
DOGAM = "원료 도감"
DOGAM_README = ("원료 도감 : 학습용 원료 766종 한 표 — 기능성·합성 130종(INCI·전형%·왜 쓰나·천연 대안·주의) + 허브누리 원료 사전 636종"
                "(등장수·대표%·범위·특징·주 사용 제형·대체재). ▼ 필터로 구분·역할별 탐색, ✓=보유. 설계 캔버스와 무관.")

# 원료 시트 12종 역할 → 아키타입 골격의 학습 역할(허브누리 분류)로 환산
TO_SKELETON = {
    "수상": ["용제"], "유상": ["유상·에몰리언트", "감촉·실리콘"], "유화": ["유화제", "가용화제"],
    "보습": ["보습"], "활성": ["활성"], "향": ["향료"], "방부": ["방부제"],
    "점증/pH": ["점증제", "pH조절"], "기타": ["항산화제"],
}
ROW_FORMULAS = {  # [@열] 단축형은 Excel이 열 때 #REF!로 깨뜨림 → 정식 구조적 참조
    "원료": {9: '=SUMIFS(입출고표[수량g],입출고표[원료],원료표[[#This Row],[원료]],입출고표[구분],"구매")',
             10: '=SUMIFS(입출고표[수량g],입출고표[원료],원료표[[#This Row],[원료]],입출고표[구분],"조정")',
             11: "=SUMIF(처방표[원료],원료표[[#This Row],[원료]],처방표[사용g])",
             12: "=원료표[[#This Row],[최초보유g]]+원료표[[#This Row],[구매g]]+원료표[[#This Row],[조정g]]"
                 "-원료표[[#This Row],[제조사용g]]"},
    "처방": {5: "=SUMIF(개발일지표[제품],처방표[[#This Row],[제품]],개발일지표[배치량g])",
             6: "=처방표[[#This Row],[퍼센트]]/100*처방표[[#This Row],[제조누계g]]"},
}
OWNED_LIST = '=INDIRECT("원료표[원료]")'  # 원료 시트에 행을 추가하면 드롭다운도 자동 확장


def _repair_master(wb) -> list[str]:
    """재고 연계 수식(#REF!·빈칸) 복구 + 입출고 원료 드롭다운을 원료표 전체로. 멱등."""
    done = []
    for sname, cols in ROW_FORMULAS.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        tbl = next(iter(ws.tables.values()), None)
        if tbl is None:
            continue
        top, bottom = (int(re.sub(r"\D", "", x)) for x in tbl.ref.split(":"))
        n = 0
        for r in range(top + 1, bottom + 1):
            if not ws.cell(r, 1).value:
                continue
            for c, f in cols.items():
                if ws.cell(r, c).value != f:
                    ws.cell(r, c, f)
                    n += 1
        if n:
            done.append(f"{sname} 수식 {n}칸")
    if "입출고" in wb.sheetnames:
        for dv in wb["입출고"].data_validations.dataValidation:
            if dv.formula1 and "원료!" in dv.formula1:
                dv.formula1 = OWNED_LIST
                done.append("입출고 원료 드롭다운→원료표 전체")
    return done


def _update_readme(wb, key: str = SHEET, line: str = README_LINE, after: str = "처방") -> None:
    """읽어보기 시트 지도에 '<key> : …' 한 줄을 <after> 줄 바로 아래 추가(있으면 문구만 갱신)."""
    if "읽어보기" not in wb.sheetnames:
        return
    ws = wb["읽어보기"]
    rows = {str(ws.cell(r, 2).value or "").split(" :")[0]: r for r in range(1, ws.max_row + 1)}
    if key in rows:  # 재실행 → 문구만 갱신
        ws.cell(rows[key], 2, line)
        return
    anchor = rows.get(after)
    if anchor is None:
        return
    r = anchor + 1
    _insert_row(ws, r)
    ws.cell(r, 1, "·").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r, 2, line).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 36
    for rr in range(1, ws.max_row + 1):  # '시트 지도 (N개)' 숫자 +1
        v = ws.cell(rr, 1).value
        m = re.match(r"(시트 지도 \()(\d+)(개\).*)", str(v or ""))
        if m:
            ws.cell(rr, 1, f"{m.group(1)}{int(m.group(2)) + 1}{m.group(3)}")
            break


def _backup(master: Path) -> Path:
    """엑셀에서 열려 있으면 중단, 아니면 타임스탬프 백업."""
    import shutil
    from datetime import datetime
    if not master.exists():
        raise SystemExit(f"없음: {master}")
    if (master.parent / f"~${master.name}").exists():
        raise SystemExit(f"⛔ 엑셀에서 '{master.name}'를 닫고 다시 실행하세요.")
    backup = master.with_name(f"{master.stem}_백업_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    shutil.copy2(master, backup)
    return backup


def _cell(ws, r, c, v=None, *, fill=None, bold=False, size=None, color=None, wrap=False,
          center=False, border=True):
    cell = ws.cell(r, c) if v is None else ws.cell(r, c, v)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if bold or size or color:
        cell.font = Font(bold=bold, size=size or 11, color=color)
    if wrap or center:
        cell.alignment = Alignment(wrap_text=wrap, vertical="top" if not center else "center",
                                   horizontal="center" if center else None)
    if border:
        cell.border = GRID
    return cell


def _navy_head(ws, r, c, text):
    cc = ws.cell(r, c, text)
    cc.font = Font(bold=True, color="FFFFFFFF")
    cc.fill = PatternFill("solid", fgColor=NAVY)
    cc.border = GRID
    cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cc


def _build_dogam(wb, owned) -> tuple[int, int]:
    """원료 도감 시트: 기능성 팔레트 + 원료 사전을 한 표로(필터 하나로 구분·역할 탐색). (원료 수, 보유 ✓ 수)."""
    if DOGAM in wb.sheetnames:
        del wb[DOGAM]
    ws = wb.create_sheet(DOGAM)
    ws.sheet_properties.tabColor = GREEN_TAB
    heads = ["원료", "역할", "전형·대표%", "구분", "등장수", "사용% 범위(최소~중앙~최대)", "기능 · 특징",
             "주 사용 제형 / 천연 대안", "주의", "대체재(모음)", "보유(원료 시트)"]
    widths = [28, 12, 11, 9, 7, 16, 44, 28, 32, 30, 20]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    span = len(heads)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    t = ws.cell(1, 1, "원료 도감 — 학습용 원료 한 표 (기능성·합성 + 허브누리 원료 사전)")
    t.font = Font(bold=True, size=14, color="FFFFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    kn = _read_knowledge(Path("원료_지식_확장.xlsx"))
    gl = _read_glossary(Path("원료_사전.xlsx"))
    fixkeys = {glossary_key(k) for k in _ROLE_FIX}
    n_fix = sum(1 for g in gl if glossary_key(g[1]) in fixkeys)
    desc = (f"설계 캔버스(처방설계)와 무관한 공부용 원료 사전. 써보고 싶은 원료는 사서 '원료' 시트에 행 추가 → 처방설계 드롭다운에 등장.   ┃ "
            f"구분: 기능성 {len(kn)}종(원료_지식_확장 — 같은 역할 안에서 '언제 이걸 고르나' 비교) · 허브누리 {len(gl)}종(원료_사전 — 레시피 통계, "
            f"역할은 이름 자동분류 + 수동 보정 {n_fix}건). ▼ 필터로 구분·역할·제형 탐색. 보유 ✓=원료 시트 이름·세이프.   ┃ "
            "⚠ 전형%·고시 한도는 업계 통용 예시 — 집행 전 식약처 고시·공급사 TDS로 검증.")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    d = ws.cell(2, 1, desc)
    d.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    d.fill = PatternFill("solid", fgColor=GREY)
    d.font = Font(size=10, color="FF404040")
    ws.row_dimensions[2].height = 46
    for j, h in enumerate(heads, 1):
        _navy_head(ws, 3, j, h)
    ws.row_dimensions[3].height = 30

    def lines(text, per):
        return max(1, -(-len(str(text or "")) // per))

    r, n_owned = 4, 0
    rows = [(name, role, pct, "기능성", None, None, why, f"대안: {alt}" if alt and alt != "—" else alt, caution, None)
            for role, name, pct, why, alt, caution in kn]
    rows += [(f"{name}\n(= {others})" if others else name, role, pct, "허브누리", cnt, rng_, feat, forms, None, alts)
             for role, name, cnt, pct, feat, alts, rng_, forms, others in gl]
    for vals in rows:
        mine = _match_owned(str(vals[0]).split("\n")[0], owned)
        n_owned += bool(mine)
        for j, v in enumerate((*vals, mine), 1):
            c = _cell(ws, r, j, v, wrap=j in (1, 6, 7, 8, 9, 10, 11))
            if j in (2, 3, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=_role_color(vals[1]))
        ws.cell(r, 4).font = Font(size=9, bold=vals[3] == "기능성", color="FF1F3A5F" if vals[3] == "기능성" else "FF55606E")
        ws.cell(r, 10).font = Font(size=9, color="FF55606E")
        if mine:
            ws.cell(r, 11).font = Font(bold=True, color="FF2E7D32")
        h = 15 * max(lines(vals[0], 14), lines(vals[6], 22), lines(vals[7], 14), lines(vals[8], 16))
        ws.row_dimensions[r].height = min(max(18, h), 75)
        r += 1
    ws.auto_filter.ref = f"A3:{get_column_letter(span)}{r - 1}"
    ws.freeze_panes = "B4"
    return len(rows), n_owned


def build_into_master(master: Path = MASTER) -> Path:
    from dev_master import ROLES  # 원료 시트 12종 역할(단일 출처)

    backup = _backup(master)

    arches = _read_archetypes(Path("카드_아키타입"))
    x14 = _x14_validations(master)
    wb = openpyxl.load_workbook(master)
    owned = _master_ingredients(wb)

    # openpyxl이 버린 x14 드롭다운(예: 입출고 원료 목록) 복원
    restored = 0
    for sname, items in x14.items():
        if sname not in wb.sheetnames:
            continue
        for f, sq in items:
            dv = DataValidation(type="list", formula1=f if f.startswith("=") else "=" + f, allow_blank=True)
            wb[sname].add_data_validation(dv)
            dv.add(sq)
            restored += 1
    repaired = _repair_master(wb)

    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = GREEN_TAB
    # A원료 B역할 C투입% D그램 E상 F세이프 G현재고g H법적한도 I경고 J메모 | K간격 | L·M 패널
    widths = (24, 14, 8, 8, 5, 6, 9, 16, 16, 18, 2, 18, 9)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    SPAN = len(widths)
    KEY = KEY_FILL.fgColor.rgb
    NOTE = NOTE_FILL.fgColor.rgb

    # 1행 제목 · 2행 설명띠(마스터 공통 머리)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=SPAN)
    t = ws.cell(1, 1, "처방설계 — 보유 원료로 내 처방 짜기")
    t.font = Font(bold=True, size=14, color="FFFFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    desc = ("집 테스트·개발용 설계 캔버스. 노란 칸(제형·배치량·원료명·투입%)만 입력 — 나머지는 원료 시트에서 자동으로 따라온다.   ┃ "
            "원료명=원료 시트 드롭다운(새 원료는 원료 시트에 행 추가 → 여기 자동 등장) · 역할·상·세이프·현재고g·법적한도=원료 시트 값 · "
            "그램=투입%×배치량 · 경고=한도초과·재고부족(이번 배치 그램 > 현재고) · 오른쪽=역할별 합계·골격 비교·자동 점검.   ┃ "
            "확정하면 '처방' 표에 [제품·원료·상·퍼센트]로 옮겨야 재고 차감이 걸린다. 새 원료를 고를 땐 '원료 도감' 시트(학습용 766종).")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=SPAN)
    c = ws.cell(2, 1, desc)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    c.fill = PatternFill("solid", fgColor=GREY)
    c.font = Font(size=10, color="FF404040")
    ws.row_dimensions[2].height = 62

    # 3~5행: 입력·골격·주의
    _cell(ws, 3, 1, "① 제형 선택", fill=KEY, bold=True)
    ws.merge_cells("B3:D3")
    for cc in (2, 3, 4):
        _cell(ws, 3, cc, fill=KEY)
    ws.merge_cells("E3:F3")
    _cell(ws, 3, 5, "② 배치량(g)", fill=KEY, bold=True)
    _cell(ws, 3, 6, fill=KEY)
    _cell(ws, 3, 7, 100, fill=KEY, bold=True)

    R0, R1 = 8, 27  # 입력 행
    rng = lambda col: f"${col}${R0}:${col}${R1}"  # noqa: E731
    B, C, D = rng("B"), rng("C"), rng("D")

    # 오른쪽 패널(L:M) 행 배치 → 그 아래로 참고 골격표
    role_top = R0
    skel_head = role_top + len(ROLES) + 1
    skel_top = skel_head + 1
    chk_head = skel_top + len(TO_SKELETON) + 1
    checks = [
        (f'=IF(SUM({C})=0,"→ 원료를 입력하세요",IF(ABS(SUM({C})-100)<0.5,"✓ 합계 100%",'
         f'"⚠ 합계 "&ROUND(SUM({C}),1)&"% (100 맞추기)"))'),
        (f'=IF(AND(SUMIF({B},"용제",{C})>0,SUM(SUMIF({B},{{"유상·에몰리언트","감촉·실리콘"}},{C}))>0,'
         f'SUM(SUMIF({B},{{"유화제","가용화제"}},{C}))=0),"⚠ 물+오일인데 유화제 없음 → 분리","")'),
        (f'=IF(AND(SUMIF({B},"용제",{C})>0,SUMIF({B},"방부제",{C})=0),"⚠ 물 있는데 방부제 없음 → 변질","")'),
        (f'=IF(COUNTIFS($F${R0}:$F${R1},"X",{C},">0")>0,"⚠ 세이프 X 원료 "&COUNTIFS($F${R0}:$F${R1},"X",{C},">0")'
         f'&"개 → 히어로(말라세지아-세이프) 계열 불가","")'),
        (f'=IF(COUNTIF($I${R0}:$I${R1},"*재고부족*")>0,"⚠ 재고 부족 "&COUNTIF($I${R0}:$I${R1},"*재고부족*")'
         f'&"개 → 배치량↓ 또는 구매","")'),
        (f'=IF(COUNTIF($I${R0}:$I${R1},"*한도초과*")>0,"⛔ 법적 한도 초과 "&COUNTIF($I${R0}:$I${R1},"*한도초과*")&"개","")'),
        (f'=IF(SUMIF({B},"향료",{C})>0,"ℹ 향료 포함 — 다도기는 무향 원칙","")'),
    ]
    ref_head = max(R1 + 4, chk_head + len(checks) + 2)
    ref_first = ref_head + 2
    ref_last = ref_first + 2 * len(arches) - 1
    A_RNG = f"$A${ref_first}:$A${ref_last}"
    B_RNG = f"$B${ref_first}:$B${ref_last}"

    _cell(ws, 4, 1, "목표 골격(참고)", bold=True, size=9, color="55606E")
    ws.merge_cells(f"B4:{get_column_letter(SPAN)}4")
    _cell(ws, 4, 2, f'=IFERROR(INDEX({B_RNG},MATCH($B$3,{A_RNG},0)),"← 제형을 선택하세요 (오른쪽 「골격 비교」에 내 처방이 같은 역할로 환산돼 나온다)")',
          wrap=True)
    ws.row_dimensions[4].height = 32
    _cell(ws, 5, 1, "배합 주의", bold=True, size=9, color="7A4A2A", fill=NOTE)
    ws.merge_cells(f"B5:{get_column_letter(SPAN)}5")
    _cell(ws, 5, 2, f'=IFERROR(INDEX({B_RNG},MATCH($B$3,{A_RNG},0)+1),"")', wrap=True, fill=NOTE)
    ws.row_dimensions[5].height = 44

    # 입력 표(7행 헤더, 8~27행)
    heads = ["원료명 ▼", "역할", "투입%", "그램", "상", "세이프", "현재고g", "법적 한도", "경고", "메모"]
    for i, h in enumerate(heads, 1):
        _navy_head(ws, 7, i, h)
    look = lambda col, r: f'=IF($A{r}="","",IFERROR(INDEX(원료표[{col}],MATCH($A{r},원료표[원료],0)),"?"))'  # noqa: E731
    for r in range(R0, R1 + 1):
        for ci in range(1, len(heads) + 1):
            _cell(ws, r, ci)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=KEY)
        ws.cell(r, 3).fill = PatternFill("solid", fgColor=KEY)
        ws.cell(r, 2, look("역할", r))
        ws.cell(r, 4, f'=IF($C{r}="","",ROUND($C{r}*$G$3/100,2))')
        ws.cell(r, 5, look("상", r))
        ws.cell(r, 6, look("세이프", r))
        ws.cell(r, 7, f'=IF($A{r}="","",IFERROR(ROUND(INDEX(원료표[현재고g],MATCH($A{r},원료표[원료],0)),1),""))')
        ws.cell(r, 8, look("법적 사용한도(고시)", r).replace(',"?"))', ',""))'))
        ws.cell(r, 9, f'=TRIM(IF(AND(ISNUMBER($C{r}),IFERROR($C{r}>VALUE(LEFT($H{r},FIND("%",$H{r})-1)),FALSE)),"⛔한도초과 ","")'
                      f'&IF(AND(ISNUMBER($D{r}),ISNUMBER($G{r})),IF($D{r}>$G{r},"⚠재고부족 "&ROUND($D{r}-$G{r},1)&"g",""),""))')
        for ci in (2, 5, 6):
            ws.cell(r, ci).alignment = Alignment(horizontal="center")
        ws.cell(r, 8).font = Font(size=9)
        ws.cell(r, 9).font = Font(size=9, color="FFC00000", bold=True)
    tr = R1 + 1
    for ci in range(1, len(heads) + 1):
        _cell(ws, tr, ci, fill=LBLUE)
    ws.cell(tr, 1, "합계").font = Font(bold=True)
    ws.cell(tr, 3, f"=SUM(C{R0}:C{R1})").font = Font(bold=True)
    ws.cell(tr, 4, f"=SUM(D{R0}:D{R1})").font = Font(bold=True)
    ws.conditional_formatting.add(f"F{R0}:F{R1}", CellIsRule(operator="equal", formula=['"X"'], fill=REDFILL))

    # 오른쪽 패널 ① 역할별 합계(원료 시트 12종)
    _navy_head(ws, 7, 12, "역할별 합계(%)")
    _navy_head(ws, 7, 13, "%")
    for i, role in enumerate(ROLES):
        rr = role_top + i
        _cell(ws, rr, 12, role, size=10)
        _cell(ws, rr, 13, f'=IF(SUMIF({B},L{rr},{C})=0,"",ROUND(SUMIF({B},L{rr},{C}),1))', center=True)
    # ② 골격 비교(학습 역할로 환산) — 위 '목표 골격'과 같은 이름으로 나란히 본다
    _navy_head(ws, skel_head, 12, "골격 비교(학습 역할)")
    _navy_head(ws, skel_head, 13, "내 %")
    for i, (srole, mroles) in enumerate(TO_SKELETON.items()):
        rr = skel_top + i
        _cell(ws, rr, 12, srole, fill=ROLE_COLORS.get(srole, "F0F0F0"), size=10)
        arr = "{" + ",".join(f'"{m}"' for m in mroles) + "}"
        _cell(ws, rr, 13, f'=IF(SUM(SUMIF({B},{arr},{C}))=0,"",ROUND(SUM(SUMIF({B},{arr},{C})),1))', center=True)
    # ③ 자동 점검
    ws.merge_cells(start_row=chk_head, start_column=12, end_row=chk_head, end_column=13)
    _cell(ws, chk_head, 12, "자동 점검", bold=True, fill=KEY)
    for i, fml in enumerate(checks, 1):
        rr = chk_head + i
        ws.merge_cells(start_row=rr, start_column=12, end_row=rr, end_column=13)
        cc = ws.cell(rr, 12, fml)
        cc.alignment = WRAP
        cc.font = Font(size=9, bold=True, color="FFC00000")
        ws.row_dimensions[rr].height = 28

    # 주의 문구(입력 표 아래)
    ur = R1 + 2
    ws.merge_cells(start_row=ur, start_column=1, end_row=ur, end_column=10)
    gc = ws.cell(ur, 1, "⚠ 설계·연습용 — 제조 전 공급사 TDS·pH·방부력 확인. 법적 한도는 원료 시트 N열(고시) 기준이며 빈칸=미기재(한도 없음 아님). "
                        "소량(50~100g)·라벨·1개월 내 사용.")
    gc.alignment = WRAP
    gc.font = Font(size=9, color="FFC00000")
    ws.row_dimensions[ur].height = 30

    # ── 참고 골격표(제형 15종): 제형행 = 골격, 다음행 = 배합 주의 ──
    RSPAN = 10
    ws.merge_cells(start_row=ref_head, start_column=1, end_row=ref_head, end_column=RSPAN)
    sc = ws.cell(ref_head, 1, f"참고 골격 — 제형 {len(arches)}종 (허브누리 레시피 통계 · 역할별 전형 % p25~p75 · 학습용 참고치)")
    sc.font = Font(bold=True)
    sc.fill = PatternFill("solid", fgColor=LBLUE)
    _cell(ws, ref_head + 1, 1, "제형", bold=True, fill=HDR_FILL.fgColor.rgb)
    ws.merge_cells(start_row=ref_head + 1, start_column=2, end_row=ref_head + 1, end_column=RSPAN)
    _cell(ws, ref_head + 1, 2, "목표 골격(전형 %) / 배합 주의", bold=True, fill=HDR_FILL.fgColor.rgb)
    r = ref_first
    for arche, skel, tips in arches:
        _cell(ws, r, 1, arche, bold=True, fill=KEY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=RSPAN)
        _cell(ws, r, 2, skel, wrap=True)
        ws.row_dimensions[r].height = 32
        _cell(ws, r + 1, 1, "  └ 배합 주의", size=9, color="7A4A2A", fill=NOTE)
        ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=RSPAN)
        _cell(ws, r + 1, 2, tips, wrap=True, fill=NOTE)
        ws.row_dimensions[r + 1].height = 46
        r += 2

    # 드롭다운: 제형(인라인 — 골격표는 2행씩이라 범위 참조 불가), 원료명(원료표 전체)
    dv_arche = DataValidation(type="list", formula1='"' + ",".join(a for a, _s, _t in arches) + '"',
                              allow_blank=True)
    ws.add_data_validation(dv_arche)
    dv_arche.add("B3")
    dv_ing = DataValidation(type="list", formula1=OWNED_LIST, allow_blank=True)
    ws.add_data_validation(dv_ing)
    dv_ing.add(f"A{R0}:A{R1}")
    ws.freeze_panes = "A8"

    # '처방' 바로 뒤에 배치 + 원료 도감(역설계 심화 뒤, 없으면 처방설계 뒤) + 읽어보기 시트 지도 갱신
    if "처방" in wb.sheetnames:
        target = wb.sheetnames.index("처방") + 1
        wb.move_sheet(SHEET, offset=target - wb.sheetnames.index(SHEET))
    n_pal, n_owned = _build_dogam(wb, owned)
    after = "역설계 심화" if "역설계 심화" in wb.sheetnames else SHEET
    wb.move_sheet(DOGAM, offset=wb.sheetnames.index(after) + 1 - wb.sheetnames.index(DOGAM))
    _update_readme(wb)
    _update_readme(wb, DOGAM, DOGAM_README, after=after)

    try:
        wb.save(master)
    except PermissionError:
        raise SystemExit(f"⛔ 저장 실패 — 엑셀에서 '{master.name}'를 닫고 다시 실행하세요.")
    print(f"✅ '{SHEET}' 시트 → {master.name} (보유 원료 {len(owned)}종 드롭다운 · 제형 {len(arches)}종 · '{DOGAM}' {n_pal}종 · 보유 ✓ {n_owned})")
    print(f"   백업: {backup.name} · x14 드롭다운 복원 {restored}개 · 복구: {' / '.join(repaired) or '없음'}")
    return master


if __name__ == "__main__":
    import sys
    if "--standalone" in sys.argv:
        build()
    else:
        build_into_master()
