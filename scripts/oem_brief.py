# -*- coding: utf-8 -*-
"""OEM/ODM 제품개발의뢰서 시트를 제품개발_마스터.xlsx에 추가.

1인·무경력 창업자가 공장(OEM/ODM) 담당자와 '정확히 같은 그림'을 그리기 위한
한 장짜리 개발 의뢰서(RFQ). PPT 발표가 아니라 이 스펙 시트가 공장이 실제로
견적·샘플을 돌리는 문서다. 보유원료(저자극 진정·장벽 라인) 기반 예시를 채워
백지공포 없이 '고쳐 쓰기'만 하면 되게 설계.

재생성(dev_master.py) 시 GENERATED_SHEETS 밖 시트라 통째 보존됨(라인 820~).
사용: .venv/Scripts/python scripts/oem_brief.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, ValueError):
        pass

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "제품개발_마스터.xlsx"
SHEET = "OEM 개발의뢰서"

# ── 팔레트 ──────────────────────────────────────────────
DEEP = "2E5A4B"      # 타이틀(딥그린)
SECT = "CFE3D8"      # 섹션 헤더
LABEL = "F0F4F2"     # 항목 라벨
TIPC = "8A8A8A"      # 팁 회색
WARN = "FBE9E7"      # 주의 핑크

WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center")
B = Font(bold=True)
WHITE_B = Font(bold=True, color="FFFFFF", size=13)
TIP = Font(italic=True, color=TIPC, size=9)
thin = Side(style="thin", color="D5DDD9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# 행 스펙: ("title",텍스트) | ("sect",텍스트) | ("f",항목,예시입력,작성팁) | ("note",텍스트) | ("tmpl",텍스트블록)
ROWS = [
    ("title", "다도기 · OEM/ODM 제품개발의뢰서 (Product Development Brief)"),
    ("note", "※ 파란 예시 = 확정 첫 제품 '시크 진정 밸런스 크림'(히어로) 기준으로 채워둠. 실제 값으로 고쳐 쓰세요. "
             "컨셉·스펙은 공유해도 되지만(견적에 필요), 마케팅·채널·매출 계획은 굳이 주지 마세요. 진짜 지킬 건 상표권('다도기' 출원)."),

    ("sect", "0. 의뢰 기본정보"),
    ("f", "브랜드 / 책임판매업자", "다도기 / 화장품책임판매업 (등록상태: ☐완료 ☐준비중)",
     "공장이 제일 먼저 묻는 것. '준비중'이라도 답할 수 있으면 신뢰도↑. 국내판매 필수 등록."),
    ("f", "담당자 / 연락처 / 작성일", "박○○ / 010-0000-0000 / 2026-__-__", "의뢰서엔 날짜·버전을 꼭 남겨 수정 추적."),
    ("f", "의뢰 유형", "ODM (공장 처방 기반, 내 설계안 전달) — ☐OEM(내 처방 보유) ☑ODM",
     "무경력·1인이면 보통 ODM(공장이 처방 보유). 내 홈배치 처방은 '설계 방향'으로 전달하고 안정성·방부는 공장 처방에 맡김. '처방 독점' 여부 꼭 협의."),
    ("f", "희망 출시 시점", "2027년 상반기 (샘플은 2개월 내 희망)", "리드타임 역산용. 무리한 일정은 오히려 불신 유발."),

    ("sect", "1. 제품 개요 (가장 중요 — 여기서 공장이 '이해'함)"),
    ("f", "제품명(가칭)", "다도기 시크 진정 밸런스 크림", "확정 아니어도 OK. 카테고리가 드러나게."),
    ("f", "카테고리 / 제형", "스킨케어 / 유화 크림(리치, 오일 약 17%)", "제형 타입은 2절에서 상세."),
    ("f", "타겟 고객(페르소나)", "20~40대, 지성·예민·트러블(지루성 경향) 피부, 유분 걱정 없이 진정·보습을 원하는 사람",
     "연령·성별·피부고민 한 줄. 구체적일수록 처방 방향이 또렷."),
    ("f", "한 줄 컨셉 / 차별점(POD)",
     "피지·예민 피부용 무향 진정+보습 밸런스 — '말라세지아 먹이'가 되는 유성성분(올리브·코코넛·호호바·버터 등 C11~24 지방산 오일)과 향료를 전면 배제",
     "경쟁사와 '무엇이 다른가' 한 문장. 이 '세이프 오일 설계'가 다도기의 핵심 정체성이자 소구점."),
    ("f", "벤치마크(레퍼런스) 2~3개", "① ○○시카크림(진정감 레퍼런스) ② △△장벽크림(보습·무향) — '이 느낌·이건 좋고/이건 싫다'",
     "★공장이 가장 빨리 이해하는 수단. 실제 시판품 실물/사진을 들고 가면 최고."),

    ("sect", "2. 제형 · 사용감 스펙"),
    ("f", "제형 타입", "유화 크림(리치) — ☐로션 ☑크림 ☐젤 ☐세럼 ☐밤 ☐오일 ☐클렌징",
     "해당에 체크 + 보충설명. 오일상 약 17% 수준(공장이 안정성에 맞게 조정 가능)."),
    ("f", "점도 / 텍스처", "리치하지만 무겁지 않게, 뭉침 없이 발리는 크림", "'물 같게/꿀 같게' 등 비유가 잘 통함."),
    ("f", "발림성·흡수·마무리", "부드럽게 밀리고 흡수 후 보송~촉촉(번들거림 적게)", "촉촉/산뜻/보송 중 끝느낌을 명시. 지성타겟이라 '번들거림 적게'가 포인트."),
    ("f", "색상 / 향", "아이보리~화이트 / ★무향(향료·에센셜오일 불가)", "무향이 민감·지루성 피부엔 안전하고, 다도기 원칙. 향 넣으면 알러젠·단가↑."),
    ("f", "pH 희망대", "약산성 5.0~6.0", "피부 약산성대. 공장이 방부·안정성과 함께 조정."),

    ("sect", "3. 성분 요구사항"),
    ("f", "꼭 넣을 핵심(소구) 성분",
     "나이아신아마이드(≈2%), 아연PCA, 마데카소사이드(시카), 베타글루칸, 판테놀, 알란토인, 저분자 히알루론산, 녹차추출물",
     "마케팅에 내세울 성분. 함량은 공장과 협의(표방·단가·갈변 안정성 영향). ★녹차·시카 갈변/안정성은 홈배치 한계였으니 공장 안정화 버전으로."),
    ("f", "세이프 에몰리언트/유화제", "스쿠알란·MCT·디메치콘 / 유화제: 글루코사이드계(예: 몬타노브68)",
     "유성성분은 '말라세지아-세이프'만. 공장 보유 유화 베이스로 대체해도 되나 '세이프 오일 유지'는 필수 조건."),
    ("f", "반드시 뺄 성분(무첨가)", "★향료·에센셜오일, C11~24 지방산 오일(올리브·코코넛·호호바·버터류), 인공색소 / (가능하면)파라벤",
     "'무첨가' 표방하려면 리스트 명확히. 단, 방부는 안전상 대체제(페녹시+헥산다이올 등) 필요."),
    ("f", "비건 / 동물유래", "비건 지향(동물유래·동물실험 배제)", "비건 인증 원하면 미리. 원료 선택·단가에 영향."),
    ("f", "기능성 여부", "일반화장품 — ☐미백 ☐주름 ☐자외선차단 ☑없음",
     "나이아신아마이드 '미백' 기능성은 가능성만 공장과 협의. 첫 제품은 심사·시험·기간·단가 때문에 일반 추천."),

    ("sect", "4. 용량 · 용기 · 부자재"),
    ("f", "내용량", "50ml", "단가·용기와 직결. 크림은 자(jar)/튜브 50ml가 무난."),
    ("f", "용기 타입", "자(jar) 또는 튜브 — ☐펌프 ☑자 ☑튜브 ☐유리 ☐에어리스",
     "크림은 자·튜브가 흔함. 자본 절감 위해 본배치는 튜브도 검토(내 설계노트 참고)."),
    ("f", "용기 조달 방식", "1차: 공장 재고(스톡) 용기 — ☑스톡(저렴) ☐전용금형(고가)",
     "★첫 제품은 스톡 용기로 금형비 절약. 전용금형은 수백만원+MOQ 큼."),
    ("f", "2차 포장 / 라벨", "단상자 미정 / 라벨·스티커 필요(전성분·바코드 포함)",
     "단상자·라벨 포함여부로 견적이 달라짐."),

    ("sect", "5. 수량 · 단가 · 일정"),
    ("f", "희망 MOQ", "가능한 낮게(예: 500~1,000개) 문의", "★공장별 MOQ 천차만별. 1인은 낮은 MOQ가 생존 핵심."),
    ("f", "1차 발주 희망 수량", "미정 / MOQ 수준부터", "과장 금지. 솔직한 소량이 오히려 신뢰."),
    ("f", "목표 소비자가 / 공급단가", "소비자가 2~3만원대 / 목표 공급단가(원가) ___원",
     "목표가를 주면 공장이 그 안에서 처방·용기를 맞춰줌."),
    ("f", "일정", "샘플 __월 / 양산 __월", "샘플→수정→양산 사이클 보통 2~4개월+."),

    ("sect", "6. 유통 · 규제"),
    ("f", "판매 채널", "자사몰 + (추후)올리브영 / 수출: ☐국내만 ☐수출(국가:____)",
     "채널에 따라 규격·표시·바코드가 달라짐."),
    ("f", "수출 대상국 규제", "예: 미국 OTC/FDA, EU CPNP·CPSR, 중국 위생허가 등(해당 시)",
     "수출이면 대상국 규제·서류를 공장이 지원하는지 확인."),
    ("f", "필요 서류", "시험성적서, 안정성·방부력(챌린지)자료, 전성분, (수출시)CPSR",
     "★공장이 이 서류를 제공하는지 반드시 확인. 없으면 위험신호."),

    ("sect", "7. 공장에 '역으로' 요청할 견적 항목 (체크리스트)"),
    ("note", "아래를 '항목별로' 받으세요. 뭉뚱그린 견적·구두 약속·과도한 선금은 경계 신호입니다."),
    ("f", "① 처방개발비", "유/무 · 금액 · 환불·전환 조건", "무료라며 유도 후 청구하는 경우 있음 — 문서로."),
    ("f", "② 샘플비 / 샘플 횟수", "건당 비용 · 무상 수정 횟수", "샘플 없이 대량발주 유도는 위험."),
    ("f", "③ 금형비(용기)", "스톡이면 0 / 전용금형 금액", "전용금형은 큰 돈·장기 묶임."),
    ("f", "④ 수량별 단가표", "500/1,000/3,000개 단가", "수량 구간별로 받아 손익분기 계산."),
    ("f", "⑤ MOQ · 리드타임", "최소수량 · 샘플/양산 소요일", ""),
    ("f", "⑥ 부자재 포함 여부", "용기·펌프·단상자·라벨 포함/별도", "'단가'에 뭐가 포함인지 꼭 확인."),
    ("f", "⑦ 처방 귀속 · 독점", "처방 소유자 / 타 브랜드 판매 제한 가능?", "★ODM 처방은 보통 공장 소유 → 독점 원하면 명시."),
    ("f", "⑧ 시험성적서 제공", "제공 여부 · 비용", "품질·안전 입증 서류."),

    ("sect", "8. 공장 첫 문의 — 이메일 템플릿 (복사해서 쓰기)"),
    ("note", "팁 — 공장 홈페이지에 '개발문의' 폼이 있으면 그걸 우선, 없으면 대표메일로. 이 의뢰서를 PDF로 내보내 첨부하세요. ○○·___는 내 값으로 교체."),
    ("tmpl",
     "제목: [다도기] 지성·예민 피부용 진정 밸런스 크림 OEM/ODM 개발 문의 (소량 발주 가능 여부)\n"
     "\n"
     "안녕하세요. 스킨케어 브랜드 '다도기'를 준비 중인 박○○입니다.\n"
     "(화장품책임판매업 등록: 완료/준비 중)\n"
     "\n"
     "지성·예민 피부를 위한 '무향 진정 밸런스 크림' 1종을 첫 제품으로 개발하려 합니다.\n"
     "상세 스펙을 정리한 개발의뢰서를 첨부드리며, 핵심만 요약하면 다음과 같습니다.\n"
     "\n"
     " · 제형   : 유화 크림(오일 약 17%), 무향, 약산성(pH 5~6)\n"
     " · 콘셉트 : 피지·예민 피부용 진정+보습 밸런스 — 특정 유성성분(C11~24 지방산 오일)·향료 배제\n"
     " · 핵심성분: 나이아신아마이드·시카(마데카소사이드)·베타글루칸·판테놀 등\n"
     " · 용량/용기: 50ml, 스톡(재고) 용기 우선\n"
     " · 채널   : 자사몰 시작 → (추후) 올리브영\n"
     "\n"
     "아래 항목으로 견적을 부탁드립니다.\n"
     " ① 처방개발비/샘플비  ② 수량별 단가(500·1,000·3,000개)  ③ MOQ·리드타임\n"
     " ④ 금형비 유무  ⑤ 부자재(용기·단상자·라벨) 포함 여부  ⑥ 처방 독점 가능 여부  ⑦ 시험성적서 제공 여부\n"
     "\n"
     "소량으로 시작하는 신규 브랜드라, 가능한 가장 낮은 MOQ 조건도 함께 알려주시면 감사하겠습니다.\n"
     "괜찮으시면 간단한 상담(전화/방문) 일정도 잡고 싶습니다.\n"
     "\n"
     "감사합니다.\n"
     "박○○ 드림  |  연락처 010-0000-0000  |  이메일 ___@___"),

    ("sect", "9. 공장 첫 문의 — 카톡/문자 짧은 버전"),
    ("tmpl",
     "안녕하세요, 스킨케어 브랜드 '다도기' 준비 중인 박○○입니다.\n"
     "지성·예민 피부용 무향 진정 크림(50ml) 1종 OEM/ODM 개발을 문의드립니다.\n"
     "신규 소량 브랜드라 ①낮은 MOQ 가능 여부 ②수량별 단가 ③샘플·개발비 ④리드타임이 궁금합니다.\n"
     "상세 개발의뢰서(PDF) 보내드릴 이메일 주소 알려주시면 바로 전달드리겠습니다.\n"
     "상담 가능한 시간도 알려주시면 감사하겠습니다. 감사합니다!"),

    ("sect", "10. 체크 — 보내기 전 / 미팅 전"),
    ("note", "□ 상표(브랜드명 '다도기') 출원했는가  □ 공장 2~3곳 비교했는가  "
             "□ 식약처 제조업 등록된 정상 공장인가  □ 책임판매업 등록(or 준비) 답변 가능한가  "
             "□ 레퍼런스 제품 실물/사진 준비했는가  □ 목표가·희망수량 정했는가  □ 이 의뢰서를 PDF로 내보냈는가"),
]


def _title(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text)
    c.font = WHITE_B
    c.fill = PatternFill("solid", fgColor=DEEP)
    c.alignment = WRAP_C
    ws.row_dimensions[r].height = 30


def _sect(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text)
    c.font = B
    c.fill = PatternFill("solid", fgColor=SECT)
    c.alignment = WRAP_C
    ws.row_dimensions[r].height = 22


def _note(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text)
    c.font = Font(italic=True, size=9, color="5A5A5A")
    c.fill = PatternFill("solid", fgColor=WARN)
    c.alignment = WRAP
    ws.row_dimensions[r].height = 34


def _tmpl(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text)
    c.font = Font(name="Consolas", size=10, color="1F3B30")
    c.fill = PatternFill("solid", fgColor="F7FAF8")
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    c.border = BORDER
    ws.row_dimensions[r].height = 15 * (text.count("\n") + 1) + 10


def _field(ws, r, label, ex, tip):
    a = ws.cell(r, 1, label); a.font = B
    a.fill = PatternFill("solid", fgColor=LABEL); a.alignment = WRAP
    b = ws.cell(r, 2, ex); b.font = Font(color="1F4E79"); b.alignment = WRAP
    c = ws.cell(r, 3, tip); c.font = TIP; c.alignment = WRAP
    for col in (1, 2, 3):
        ws.cell(r, col).border = BORDER
    ws.row_dimensions[r].height = 30


README_SECTION = "창업·OEM 실무 시트 (스터디 밖)"
README_ITEMS = [
    "OEM 개발의뢰서 : 공장에 보낼 '제품개발의뢰서(RFQ)' 한 장. 첫 제품=히어로 '시크 진정 밸런스 크림' 예시가 채워져 있다 — 고쳐 쓰기만.",
    "공장엔 PPT 발표가 아니라 이 스펙 시트를 PDF로 보낸다. 레퍼런스 제품 실물/사진을 같이 준비하면 전달이 가장 정확.",
    "줘도 되는 것 = 제형·성분·용량 등 스펙(견적에 필요) / 아껴야 할 것 = 마케팅·채널·매출 계획. 진짜 해자는 상표권('다도기' 출원).",
    "견적은 항목별로(개발비·샘플비·수량별 단가·MOQ·리드타임·처방 독점·시험성적서). 뭉뚱그린 견적·구두 약속·과도한 선금은 경계 신호.",
    "시트 아래쪽에 공장 첫 문의용 이메일·카톡 템플릿 포함 — 복사해서 바로 쓴다.",
    "브랜드 기획서 : '남에게 다도기가 뭔지' 설명하는 한 장(자사몰·입점·투자 공통 뼈대). PPT 필요하면 이 시트를 슬라이드로 옮긴다. (scripts/brand_plan.py)",
]
LBLUE = "FFDCE6F1"
CEN = Alignment(horizontal="center", vertical="center")


def _update_readme(wb) -> bool:
    """읽어보기 시트 맨 아래에 창업·OEM 안내 섹션을 1회 추가(중복 방지)."""
    if "읽어보기" not in wb.sheetnames:
        return False
    ws = wb["읽어보기"]
    # 이미 있으면 그 섹션(앞 빈 줄 포함)을 지우고 최신 내용으로 다시 씀
    start = next((row[0].row for row in ws.iter_rows(min_col=1, max_col=1)
                  if row[0].value == README_SECTION), None)
    if start is not None:
        gap = 1 if start > 1 and ws.cell(start - 1, 1).value is None else 0
        ws.delete_rows(start - gap, ws.max_row - (start - gap) + 1)
    r = ws.max_row + 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(r, 1, README_SECTION)
    c.font = B
    c.fill = PatternFill("solid", fgColor=LBLUE)
    c.alignment = WRAP
    r += 1
    for it in README_ITEMS:
        ws.cell(r, 1, "·").alignment = CEN
        cc = ws.cell(r, 2, it)
        cc.alignment = WRAP
        ws.row_dimensions[r].height = 15 * (len(it) // 44 + 1) + 6
        r += 1
    return True


def main() -> None:
    if not XLSX.exists():
        sys.exit(f"없음: {XLSX}")
    wb = openpyxl.load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 52
    ws.sheet_view.showGridLines = False

    r = 1
    for row in ROWS:
        kind = row[0]
        if kind == "title":
            _title(ws, r, row[1])
        elif kind == "sect":
            _sect(ws, r, row[1])
        elif kind == "note":
            _note(ws, r, row[1])
        elif kind == "f":
            _field(ws, r, row[1], row[2], row[3])
        elif kind == "tmpl":
            _tmpl(ws, r, row[1])
        r += 1

    # 컬럼 머리글(섹션1 바로 위 느낌) — 2행에 안내가 있으니 생략, 대신 틀고정
    ws.freeze_panes = "A4"

    # 유통안전·OEM검수 옆으로 배치
    try:
        idx = wb.sheetnames.index("유통안전·OEM검수") + 1
        wb.move_sheet(SHEET, offset=idx - wb.sheetnames.index(SHEET))
    except ValueError:
        pass

    readme_added = _update_readme(wb)

    wb.save(XLSX)
    print(f"✅ '{SHEET}' 시트 추가 완료 → {XLSX.name}")
    print("   열: A=항목 / B=내 입력(예시) / C=작성 팁·공장이 보는 포인트")
    print(f"   읽어보기 안내 섹션: {'추가/갱신함' if readme_added else '건너뜀'}")


if __name__ == "__main__":
    main()
