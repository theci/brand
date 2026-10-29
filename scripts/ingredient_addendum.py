# -*- coding: utf-8 -*-
"""원료 추가노트 — 제본한 1권(원료 마스터) 뒤에 쌓는 증보 워크북.

실제 제품 전성분 사진을 해부하면서 1권에 없던 원료를 신규 카드(N번호)로 쌓고,
1권 원료 247종은 역할별 색인(+1권 쪽 번호)으로 다시 묶어 '어디 있나'를 바로 찾게 한다.

시트:
  1) 읽어보기      쌓는 법 · 번호 체계(#=1권, N=신규, P=제품)
  2) 원료 색인     1권 247 + 신규 N, 역할별 묶음 · 전성분 표기(별칭) · 찾는 곳(1권 쪽) · 등장 제품
  3) 신규 원료 카드 N01~ (정체·감각·%·왜 넣나·주의·말라세지아·1권 연결) + 내 메모(입력)
  4) 제품 목록     P001~ 한 줄 요약(구조·향·배울 점) + 내 메모(입력)
  5) 제품 해부     제품별 전성분 순서대로 → 카드 번호·찾는 곳·한 줄 메모

데이터: data/ingredient_addendum.yaml (제품·신규 카드·별칭을 여기에 쌓는다)
1권 쪽 번호: 인쇄본/1권_원료마스터.pdf 의 '원료 색인'에서 읽음(pdftotext 필요, 없으면 빈칸)
산출: 원료_추가노트.xlsx (리포 루트)
사용: uv run python scripts/ingredient_addendum.py
재생성 시 '내 메모' 칸은 번호 기준으로 보존한다.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, ValueError):
        pass

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "원료_추가노트.xlsx"
MASTER = ROOT / "원료_마스터.xlsx"
BOOK1 = ROOT / "인쇄본" / "1권_원료마스터.pdf"
DATA = ROOT / "data" / "ingredient_addendum.yaml"

FONT = "맑은 고딕"
NAVY = "FF1F3864"
HEAD = "FF2F5597"
ROLE_ORDER = [
    "용제", "보습", "점증", "유화", "에몰리언트", "경화", "감촉·실리콘", "활성", "항산화",
    "방부", "pH조절", "킬레이트", "가용화", "계면", "분말", "색소", "자외선", "향", "기타", "부산물",
]
GROUP_FILL = {
    "용제": "FFDDEBF7", "보습": "FFDDEBF7", "점증": "FFE2EFDA", "유화": "FFFCE4D6",
    "에몰리언트": "FFFFF2CC", "감촉·실리콘": "FFEDEDED", "활성": "FFE4DFEC", "항산화": "FFFDE9D9",
    "방부": "FFF8CBAD", "pH조절": "FFF2F2F2", "킬레이트": "FFF2F2F2", "가용화": "FFDEEAF1",
    "분말": "FFEFEFEF", "자외선": "FFFFE699", "향": "FFF4E1F5",
    "경화": "FFFFE5B4", "계면": "FFD9F2F2", "색소": "FFFADADD", "기타": "FFF2F2F2", "부산물": "FFF2F2F2",
}
NEW_FILL = PatternFill("solid", fgColor="FFFFF2CC")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF9C4")
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(bold=False, size=10, color="FF000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)


def title(ws, text, sub, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(1, 1, text)
    c.font = f(True, 14, "FFFFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    c = ws.cell(2, 1, sub)
    c.font = f(size=9, color="FF404040")
    c.alignment = WRAP
    ws.row_dimensions[2].height = 44


def header(ws, row, cols):
    for j, (n, w) in enumerate(cols, 1):
        c = ws.cell(row, j, n)
        c.font = f(True, 10, "FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w


def put(ws, r, values, fill=None, bold=False):
    for j, v in enumerate(values, 1):
        c = ws.cell(r, j, v)
        c.font = f(bold)
        c.alignment = WRAP
        c.border = BORDER
        if fill:
            c.fill = fill


# ── 입력 ────────────────────────────────────────────────────────────────────
def load_book1() -> list[dict]:
    ws = openpyxl.load_workbook(MASTER, read_only=True)["원료 카드"]
    cards = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None or r[1] is None:
            continue
        cards.append({"no": int(r[0]), "name": str(r[1]), "inci": str(r[2] or ""),
                      "own": str(r[3] or ""), "role": str(r[4] or ""), "sub": str(r[5] or "")})
    return cards


def book1_pages() -> dict[str, int]:
    """1권 PDF 끝의 '원료 색인'(가나다순 · 이름 ..... 쪽)에서 원료명→쪽."""
    if not BOOK1.exists() or not shutil.which("pdftotext"):
        return {}
    txt = subprocess.run(["pdftotext", "-layout", str(BOOK1), "-"], capture_output=True,
                         text=True, check=True).stdout
    i = txt.find("가나다순. 오른쪽 숫자는 쪽")
    pages = {}
    for line in txt[i:].splitlines():
        for seg in re.split(r"\s{3,}", line):
            m = re.match(r"^(.+?)\s*\.{3,}\s*(\d+)$", seg.strip())
            if m:
                pages[m.group(1).strip()] = int(m.group(2))
    return pages


def old_memos() -> dict[tuple[str, str], str]:
    """기존 산출물의 '내 메모'를 (시트, 번호) 키로 보존."""
    if not OUT.exists():
        return {}
    wb = openpyxl.load_workbook(OUT)
    memos = {}
    for sheet in ("신규 원료 카드", "제품 목록"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        heads = [c.value for c in ws[4]]
        if "내 메모" not in heads:
            continue
        mi = heads.index("내 메모")
        for row in ws.iter_rows(min_row=5, values_only=True):
            if row[0] and row[mi]:
                memos[(sheet, str(row[0]))] = row[mi]
    return memos


def parse_ing(line: str) -> tuple[str, str, str]:
    parts = [p.strip() for p in str(line).split("|")]
    return parts[0], parts[1] if len(parts) > 1 else "", parts[2] if len(parts) > 2 else ""


# ── 시트 ────────────────────────────────────────────────────────────────────
def sheet_readme(wb, n_book1, n_new, n_prod):
    ws = wb.active
    ws.title = "읽어보기"
    title(ws, "원료 추가노트 — 1·2권 증보",
          "제본한 1권 뒤에 쌓는 노트. 실제 제품 전성분을 해부하며 1권에 없던 원료를 신규 카드로 추가한다. "
          "재생성: uv run python scripts/ingredient_addendum.py (데이터 = data/ingredient_addendum.yaml, 내 메모 칸은 보존).", 2)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    rows = [
        ("현재", f"1권 원료 {n_book1}종 · 신규 카드 {n_new}종 · 해부한 제품 {n_prod}개"),
        ("번호 체계", "#13 = 1권 원료 카드 No.13 (책에서 찾기) · N01 = 이 노트의 신규 카드 · P001 = 해부한 제품"),
        ("원료 색인", "역할(카테고리)별로 1권 + 신규 원료를 한 표에. '찾는 곳'에 1권 쪽 번호. "
                  "'전성분 표기'는 제품 라벨에 실제로 적히는 이름(예: 구연산 → 시트릭애씨드)."),
        ("신규 원료 카드", "1권에 없던 원료. 1권 카드보다 짧게 — 정체·감각·%·왜 넣나·주의·말라세지아·1권 연결."),
        ("제품 목록 / 해부", "사진 찍은 제품별 구조·향·배울 점, 전성분 순서대로 각 성분이 어느 카드인지."),
        ("출시 비용·절차", "등록·자격·인증·품질검사·OEM·예산을 출처로 재검증한 표(확인/정정/미확인). 2권 '개발→창업 워크플로우'의 증보."),
        ("정오표", "제본한 1·2권에서 고치거나 보충할 내용을 쪽 번호와 함께."),
        ("쌓는 루프", "① 제품 전성분 사진 → Claude에게 질문  ② 1권에 있으면 #번호, 없으면 N번호 카드 추가  "
                  "③ 라벨 표기가 1권 이름과 다르면 별칭 추가  ④ 재생성  ⑤ 손으로 '내 메모'"),
        ("전성분 읽기 요령", "1% 초과 성분은 많은 순서, 1% 이하는 순서 무관. 함량(ppm) 표기 = 마케팅 강조 성분. "
                       "원료 1개가 성분 여러 개로(부산물·용매) 적히기도 한다."),
        ("⚠", "%·규제 수치는 교육용 정리(추정 포함). 제조·판매 전 식약처 최신 고시·공급사 TDS로 재확인."),
    ]
    for i, (k, v) in enumerate(rows, 4):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.font, b.font = f(True), f()
        a.alignment = b.alignment = WRAP
        a.border = b.border = BORDER
        a.fill = PatternFill("solid", fgColor="FFD9E1F2")


def sheet_index(wb, book1, pages, aliases, new_cards, seen):
    ws = wb.create_sheet("원료 색인")
    cols = [("카테고리", 11), ("번호", 7), ("원료", 26), ("전성분 표기(별칭)", 26), ("INCI", 34),
            ("보유", 6), ("보조 역할", 24), ("찾는 곳", 14), ("등장 제품", 12)]
    title(ws, "원료 색인 — 역할별 (1권 + 신규)",
          "역할(카테고리) 묶음 안에서 1권 카드 → 신규 카드(노란 행) 순. 찾는 곳: '1권 p.18' = 제본 1권 18쪽, "
          "'신규 카드' = 이 파일 '신규 원료 카드' 시트. 등장 제품 = 해부한 제품 번호(제품 해부 시트).", len(cols))
    header(ws, 4, cols)
    rows: dict[str, list] = {}
    for c in book1:
        key = f"#{c['no']}"
        p = pages.get(c["name"])
        rows.setdefault(c["role"], []).append(
            ([c["role"], key, c["name"], ", ".join(aliases.get(c["no"], [])), c["inci"], c["own"], c["sub"],
              f"1권 p.{p}" if p else "1권", ", ".join(seen.get(key, []))], None))
    for c in new_cards:
        rows.setdefault(c["role"], []).append(
            ([c["role"], c["id"], c["name"], "", c["inci"], "", c.get("sub", ""), "신규 카드",
              ", ".join(seen.get(c["id"], []))], NEW_FILL))
    r = 5
    order = ROLE_ORDER + [k for k in rows if k not in ROLE_ORDER]
    for role in order:
        if role not in rows:
            continue
        n_new = sum(1 for _, fl in rows[role] if fl)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        c = ws.cell(r, 1, f"{role}  ·  {len(rows[role])}종" + (f" (신규 {n_new})" if n_new else ""))
        c.font = f(True, 11)
        c.fill = PatternFill("solid", fgColor=GROUP_FILL.get(role, "FFFFFFFF"))
        r += 1
        for vals, fill in rows[role]:
            put(ws, r, vals, fill)
            r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{r - 1}"


def sheet_new_cards(wb, new_cards, book1_by_no, pages, memos):
    ws = wb.create_sheet("신규 원료 카드")
    cols = [("번호", 6), ("원료", 22), ("INCI", 28), ("역할", 9), ("보조 역할", 20), ("정체 — 무엇·기원", 40),
            ("감각·사용감", 22), ("전형 %", 16), ("왜 넣나 (처방 포인트)", 44), ("주의·규제", 36),
            ("말라세지아", 16), ("1권 연결", 26), ("처음 본 제품", 9), ("내 메모", 30)]
    title(ws, f"신규 원료 카드 — 1권에 없던 원료 {len(new_cards)}종",
          "1행 = 원료 1개. '1권 연결'의 #번호는 1권 원료 카드(쪽 번호 병기). 말라세지아: ○ 우려 낮음 · △ 확인 필요 · ✕ 피함. "
          "노란 '내 메모' 칸만 내가 적는 칸 — 재생성해도 보존.", len(cols))
    header(ws, 4, cols)

    def link(text):  # "#76 마카다미아넛오일" → "#76(p.49) 마카다미아넛오일"
        def rep(m):
            c = book1_by_no.get(int(m.group(1)))
            p = pages.get(c["name"]) if c else None
            return f"#{m.group(1)}(p.{p})" if p else m.group(0)
        return re.sub(r"#(\d+)", rep, str(text or ""))

    for i, c in enumerate(new_cards, 5):
        vals = [c["id"], c["name"], c["inci"], c["role"], c.get("sub", ""), c.get("what", ""),
                c.get("feel", ""), c.get("pct", ""), c.get("why", ""), c.get("caution", ""),
                c.get("malassezia", ""), link(c.get("related", "")), c.get("first_seen", ""),
                memos.get(("신규 원료 카드", c["id"]), "")]
        put(ws, i, vals)
        ws.cell(i, 1).font = ws.cell(i, 2).font = f(True)
        ws.cell(i, 4).fill = PatternFill("solid", fgColor=GROUP_FILL.get(c["role"], "FFFFFFFF"))
        ws.cell(i, len(cols)).fill = INPUT_FILL
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(new_cards) + 4}"


def resolve(ref, book1_by_no, new_by_id, pages):
    """참조(#13 / N01) → (판정, 카드 원료명, 역할, 찾는 곳)."""
    if ref.startswith("#"):
        c = book1_by_no.get(int(ref[1:]))
        if c:
            p = pages.get(c["name"])
            return "1권", c["name"], c["role"], f"1권 p.{p}" if p else "1권"
    elif ref in new_by_id:
        c = new_by_id[ref]
        return ("부산물" if c["role"] == "부산물" else "신규"), c["name"], c["role"], "신규 카드"
    return "미확인", "", "", ""


def sheet_products(wb, products, book1_by_no, new_by_id, pages, memos):
    ws = wb.create_sheet("제품 목록")
    cols = [("번호", 7), ("제품", 30), ("유형", 14), ("용량", 9), ("날짜", 11), ("성분 수", 7),
            ("1권", 6), ("신규", 6), ("핵심 구조", 48), ("향", 40), ("배울 점", 56), ("사진", 18), ("내 메모", 30)]
    title(ws, "제품 목록 — 사진으로 해부한 실제 제품",
          "1권/신규 = 전성분 중 1권 카드로 찾은 수 / 이 노트에 신규로 추가한 수(부산물 포함). 상세는 '제품 해부' 시트.", len(cols))
    header(ws, 4, cols)
    for i, p in enumerate(products, 5):
        kinds = [resolve(parse_ing(x)[1], book1_by_no, new_by_id, pages)[0] for x in p["ingredients"]]
        vals = [p["id"], p["name"], p.get("type", ""), p.get("volume", ""), str(p.get("date", "")),
                len(kinds), kinds.count("1권"), len(kinds) - kinds.count("1권"),
                p.get("structure", ""), p.get("scent", ""), p.get("lesson", ""), p.get("photo", ""),
                memos.get(("제품 목록", p["id"]), "")]
        put(ws, i, vals)
        ws.cell(i, len(cols)).fill = INPUT_FILL
    ws.freeze_panes = "C5"


def sheet_teardown(wb, products, book1_by_no, new_by_id, pages):
    ws = wb.create_sheet("제품 해부")
    cols = [("순번", 6), ("전성분 (라벨 표기)", 30), ("판정", 7), ("카드", 7), ("카드 원료명", 24),
            ("역할", 10), ("찾는 곳", 12), ("한 줄 메모", 34)]
    title(ws, "제품 해부 — 전성분을 라벨 순서대로 카드에 연결",
          "판정: 1권 = 제본 1권에 카드 있음 · 신규 = 이 노트 신규 카드(노란 행) · 부산물 = 일부러 넣지 않은 동반 성분. "
          "1% 이하 구간은 순서가 의미 없음.", len(cols))
    r = 4
    for p in products:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        c = ws.cell(r, 1, f"{p['id']}  {p['name']}  ·  {p.get('type', '')}  ·  {p.get('volume', '')}")
        c.font = f(True, 12, "FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        r += 1
        for label in ("structure", "scent", "note"):
            if p.get(label):
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
                name = {"structure": "구조", "scent": "향", "note": "메모"}[label]
                c = ws.cell(r, 1, f"{name}: {p[label]}")
                c.font = f(size=9)
                c.alignment = WRAP
                ws.row_dimensions[r].height = 30
                r += 1
        header(ws, r, cols)
        r += 1
        for n, line in enumerate(p["ingredients"], 1):
            label, ref, memo = parse_ing(line)
            kind, cname, role, where = resolve(ref, book1_by_no, new_by_id, pages)
            put(ws, r, [n, label, kind, ref, cname, role, where, memo], NEW_FILL if kind in ("신규", "부산물") else None)
            ws.cell(r, 6).fill = PatternFill("solid", fgColor=GROUP_FILL.get(role, "FFFFFFFF"))
            r += 1
        r += 1


VERDICT_FILL = {"확인": "FFE2EFDA", "정정": "FFF8CBAD", "미확인": "FFFFF2CC", "보충": "FFDDEBF7"}


def sheet_launch(wb, launch):
    ws = wb.create_sheet("출시 비용·절차")
    cols = [("구분", 14), ("항목", 22), ("확인된 내용", 70), ("금액(추정)", 18), ("판정", 8), ("출처", 40)]
    title(ws, "출시 비용·절차 — 조사 재검증 (2026-10)",
          launch.get("note", "") + "  판정: 확인 = 출처로 맞음 · 정정 = 조사 내용이 틀려 고침 · 미확인 = 공개 근거 부족, 견적·문의로 확정.",
          len(cols))
    header(ws, 4, cols)
    src = launch.get("sources", {})
    r, prev = 5, None
    for g, item, body, amount, verdict, key in launch["rows"]:
        if g != prev:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
            c = ws.cell(r, 1, g)
            c.font = f(True, 11)
            c.fill = PatternFill("solid", fgColor="FFD9E1F2")
            r, prev = r + 1, g
        put(ws, r, [g, item, body, amount, verdict, src.get(key, "")])
        ws.cell(r, 2).font = f(True)
        ws.cell(r, 5).fill = PatternFill("solid", fgColor=VERDICT_FILL.get(verdict, "FFFFFFFF"))
        ws.cell(r, 5).alignment = CENTER
        if src.get(key):
            ws.cell(r, 6).hyperlink = src[key]
            ws.cell(r, 6).font = f(color="FF0563C1")
        r += 1
    ws.freeze_panes = "C5"


def sheet_errata(wb, errata, sources):
    ws = wb.create_sheet("정오표")
    cols = [("권", 6), ("쪽", 14), ("시트", 26), ("책에 적힌 것", 36), ("정정·보충", 60), ("종류", 7), ("근거", 40)]
    title(ws, "정오표 — 제본한 1·2권에 덧붙일 정정·보충",
          "쪽 = 인쇄본 PDF 쪽 번호. 정정 = 책 내용이 틀림 · 보충 = 맞지만 빠진 정보. 책에 포스트잇으로 붙이거나 이 표를 출력해 끼워 두기.",
          len(cols))
    header(ws, 4, cols)
    for i, (vol, page, sheet, old, new, kind, key) in enumerate(errata, 5):
        put(ws, i, [vol, page, sheet, old, new, kind, sources.get(key, "")])
        ws.cell(i, 6).fill = PatternFill("solid", fgColor=VERDICT_FILL.get(kind, "FFFFFFFF"))
        if sources.get(key):
            ws.cell(i, 7).hyperlink = sources[key]
            ws.cell(i, 7).font = f(color="FF0563C1")


def main():
    data = yaml.safe_load(DATA.read_text(encoding="utf-8"))
    book1 = load_book1()
    book1_by_no = {c["no"]: c for c in book1}
    new_cards = data.get("new_cards", [])
    new_by_id = {c["id"]: c for c in new_cards}
    products = data.get("products", [])
    aliases = {int(k): v for k, v in (data.get("aliases") or {}).items()}
    pages = book1_pages()
    memos = old_memos()

    seen: dict[str, list[str]] = {}
    for p in products:
        for line in p["ingredients"]:
            ref = parse_ing(line)[1]
            if ref and p["id"] not in seen.setdefault(ref, []):
                seen[ref].append(p["id"])
    bad = [(p["id"], parse_ing(x)) for p in products for x in p["ingredients"]
           if resolve(parse_ing(x)[1], book1_by_no, new_by_id, pages)[0] == "미확인"]

    wb = openpyxl.Workbook()
    sheet_readme(wb, len(book1), len(new_cards), len(products))
    sheet_index(wb, book1, pages, aliases, new_cards, seen)
    sheet_new_cards(wb, new_cards, book1_by_no, pages, memos)
    sheet_products(wb, products, book1_by_no, new_by_id, pages, memos)
    sheet_teardown(wb, products, book1_by_no, new_by_id, pages)
    launch = data.get("launch")
    if launch:
        sheet_launch(wb, launch)
    if data.get("errata"):
        sheet_errata(wb, data["errata"], (launch or {}).get("sources", {}))
    wb.save(OUT)
    print(f"저장: {OUT.name}  (1권 {len(book1)} · 쪽번호 {len(pages)} · 신규 {len(new_cards)} · 제품 {len(products)})")
    for pid, ing in bad:
        print(f"  ⚠ 미확인 참조 {pid}: {ing}")


if __name__ == "__main__":
    main()
