# -*- coding: utf-8 -*-
"""원료 지식 확장: 기능성·합성 원료 팔레트 + 비율·제형 민감도 사다리.

허브누리 레시피는 천연 DIY 위주라, 성능을 더 올릴 때 쓰는 기능성 원료 지식이
비어 있다. 이 스크립트는 그 공백을 메우는 학습 레퍼런스를 만든다(아키타입
카드·원료사전과 짝).

⚠ 성격: 업계 통용 배합지식(예시). 전형%·고시·안전성·방부 적정성은 집행 전
식약처 고시·공급사 TDS·방부력(챌린지) 테스트로 반드시 검증할 것.

사용: uv run python scripts/ingredient_knowledge.py
산출: 원료_지식_확장.xlsx (시트 2개)
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 역할 색상(herbnoori_detail.ROLE_COLORS와 통일).
ROLE_COLORS = {
    "유화": "D8F5D8", "방부": "E2E5E9", "점증/pH": "E9F0D6", "활성": "EAD9F7",
    "에몰리언트": "FFF0CC", "보습": "CFF3F0", "항산화": "FBE9D0", "기타": "F0F0F0",
}

# ── A) 기능성·합성 원료 팔레트 ────────────────────────────────────────────
# (역할, 원료[INCI], 전형%, 기능/왜 쓰나, 천연 대안, 주의)
PALETTE = [
    # 유화제 ----------------------------------------------------------------
    ("유화", "몬타노브68 (Cetearyl Glucoside & Cetearyl Alcohol)", "3~5%",
     "글루코사이드계 O/W 유화제. 액정(라멜라) 구조로 장벽 유사·저자극. 히어로가 사용.",
     "올리브유화왁스", "글루코사이드계라 전해질·저pH에 비교적 강함. 공유화 세틸알콜로 점도↑"),
    ("유화", "올리벰1000 (Cetearyl Olivate & Sorbitan Olivate)", "3~6%",
     "올리브 유래 O/W 유화제. 액정형성·에몰리언트감. 저자극 범용.", "올리브유화왁스",
     "고온(70℃↑) 유화, 냉각 교반으로 점도 안정"),
    ("유화", "이멀시파잉 왁스 NF / 폴라왁스 (Cetearyl Alcohol & Polysorbate60)", "3~6%",
     "범용 O/W 유화제. 안정적·실패 적어 입문용.", "—(합성계)",
     "PEG계 포함 — ‘천연’ 포지션과 상충될 수 있음"),
    ("유화", "GMS SE (Glyceryl Stearate SE)", "2~4%",
     "자기유화형 보조 유화제·점도부여. 단독보단 메인유화제와 병용.", "—",
     "SE=self-emulsifying. 단독 사용 시 안정성 약함"),
    ("유화", "세테아릴알코올 (Cetearyl Alcohol)", "2~5%",
     "지방알코올. 공유화+점도+에몰리언트(3역). 크림 바디감 핵심.", "—",
     "유화제 아님(보조). 과다 시 왁시·답답함"),
    ("유화", "레시틴 (Lecithin)", "0.5~3%",
     "천연 인지질 유화제(W/O 경향). 침투·에몰리언트.", "자체가 천연",
     "산패 쉬움(항산화 병용). 끈적·어두운 색"),
    ("유화", "폴리소르베이트20/80 (Polysorbate)", "1~3%",
     "가용화제. 소량 향·오일을 물에 투명 분산(미스트·토너).", "올리브리퀴드",
     "세정력 약간. EO:가용화제 대략 1:3~1:5"),
    # 방부 ------------------------------------------------------------------
    ("방부", "1,2-헥산디올 (1,2-Hexanediol)", "2~5%",
     "다기능 습윤+방부보조. ‘무방부’ 표방 제형의 베이스.", "—",
     "단독 방부력 부족 — 카프릴릴글라이콜 등과 병용 권장"),
    ("방부", "카프릴릴글라이콜+에틸헥실글리세린", "0.3~1%",
     "다기능 방부보조+에몰리언트. 1,2-헥산디올과 시너지.", "—", "단독 완전방부 아님"),
    ("방부", "페녹시에탄올(+에틸헥실글리세린)", "0.5~1%",
     "광범위 방부. 가장 보편적 메인 방부.", "—",
     "화장품 배합한도 1% — 초과 금지. ‘천연’ 포지션과 상충"),
    ("방부", "포타슘소르베이트+소듐벤조에이트", "각 0.1~0.5%",
     "식품등급 방부. 저pH(5.0 이하)에서만 효력.", "자체가 천연유래",
     "pH 5 초과면 방부력 급감 — 반드시 pH 조정"),
    ("방부", "류시덜(Leuconostoc/무발효물)", "2~4%",
     "발효 유래 천연 방부. 천연 포지션에 적합.", "자체가 천연",
     "효력 약·배치편차 — 챌린지테스트 필수, 단기·냉장 고려"),
    # 점증/pH ---------------------------------------------------------------
    ("점증/pH", "잔탄검 (Xanthan Gum)", "0.2~1%",
     "천연 다당 점증. 가벼운 젤·유화 안정. 글리세린에 선분산.", "자체가 천연",
     ">0.7% 실뭉침·늘어짐. 전해질에 점도 변동"),
    ("점증/pH", "카보머 (Carbomer) + 중화제", "0.1~0.5%",
     "투명·맑은 젤. 중화(TEA/NaOH)로 젤화. 하이엔드 질감.", "잔탄·스클레로튬검",
     "산성·전해질·고농도알콜에 무너짐. 중화 필수"),
    ("점증/pH", "하이드록시에틸셀룰로오스 (HEC)", "0.5~1.5%",
     "셀룰로오스 점증. 맑고 늘어지지 않는 점도.", "잔탄",
     "수화에 시간·교반 필요(덩어리 주의)"),
    ("점증/pH", "구연산 / 소듐시트레이트", "0.05~0.5%",
     "pH 다운/버퍼. 약산성(5~6) 맞춤·방부 효력 확보.", "자체가 천연", "소량씩 적정하며 pH 측정"),
    ("점증/pH", "트리에탄올아민(TEA)/소듐하이드록사이드", "미량",
     "pH 업·카보머 중화제.", "—", "강염기 — 취급 주의, 과량 시 자극"),
    ("점증/pH", "소듐글루코네이트 / EDTA (킬레이트)", "0.1~0.2%",
     "금속이온 봉쇄 → 방부·산패·변색 방지, 세정력 보조.", "소듐파이테이트(천연)", "소량으로 충분"),
    # 활성 ------------------------------------------------------------------
    ("활성", "나이아신아마이드 (Niacinamide, B3)", "2~5%",
     "식약처 고시 미백. 장벽강화·피지조절·모공·톤. 가성비 최상.", "—",
     ">5% 또는 저pH서 홍조 가능. 순수VitC와 동시 고농도 주의"),
    ("활성", "비타민C 유도체 (Ascorbyl Glucoside/SAP/EAC)", "0.5~2%",
     "안정형 VitC. 미백·항산화. 순수 아스코르브산보다 안정.", "—",
     "유도체별 안정 pH·전환율 상이 — TDS 확인"),
    ("활성", "레티놀 / 레티날 (Retinol)", "0.1~1%(캡슐)",
     "턴오버·주름 개선 강력 활성.", "바쿠치올(식물성 대안)",
     "광·산소 불안정, 자극·임부 주의. 캡슐·저농도 시작"),
    ("활성", "펩타이드 (Matrixyl 등 Palmitoyl peptides)", "2~5%(원료)",
     "콜라겐 신호·주름. 저자극.", "—", "실활성은 ppm 수준 — 원료 투입%와 혼동 주의, 열 피함"),
    ("활성", "판테놀 (Panthenol, B5)", "1~5%",
     "보습·진정·장벽. 무난한 만능 활성.", "—", "거의 무자극, 고함량도 안전한 편"),
    ("활성", "알란토인 (Allantoin)", "0.1~0.5%",
     "진정·각질유연. 민감·트러블.", "컴프리추출물", "용해도 낮음(0.5% 내외 한계)"),
    ("활성", "아데노신 (Adenosine)", "0.04%",
     "식약처 고시 주름개선 농도.", "—", "고시농도 0.04% 준수"),
    ("활성", "살리실산 (BHA, Salicylic Acid)", "0.5~2%",
     "지용성 각질·모공·피지(지성·트러블).", "버드나무껍질추출(천연 BHA)",
     "저pH서 효력·자극. 씻어내기/바르기 농도 규정 상이·임부 주의"),
    ("활성", "아연 PCA / 나이아신 (피지조절)", "0.1~1%",
     "피지·유분 조절(지성·지루 경향). 히어로 결에 부합.", "—", "과건조 주의"),
    # 에몰리언트 ------------------------------------------------------------
    ("에몰리언트", "스쿠알란 (Squalane)", "2~15%",
     "초경량·안정 에몰리언트. 말라세지아-세이프(히어로 사용).", "올리브스쿠알란(식물)",
     "말라세지아 먹이(C11~24 지방산) 아님 — 지루·예민에 안전"),
    ("에몰리언트", "카프릴릭/카프릭트라이글리세라이드 (MCT, C8/C10)", "3~15%",
     "가볍고 산뜻·안정. 말라세지아-세이프.", "분획 코코넛오일",
     "C8/C10이라 효모 먹이 아님 — 지성·예민 적합"),
    ("에몰리언트", "다이메치콘/사이클로메치콘 (Silicone)", "1~10%",
     "실키·발림·막형성. 비유성 매끈함. 모공 비유발 경향.", "—(실리콘대안: 스쿠알란)",
     "‘천연’ 포지션과 상충 가능. 사이클로는 휘발"),
    ("에몰리언트", "C12-15 알킬벤조에이트", "2~10%",
     "초경량 에스터. 산뜻·안정·향 용해.", "—", "합성 에스터 — 포지션 고려"),
    ("에몰리언트", "아이소프로필미리스테이트(IPM)/팔미테이트", "2~8%",
     "가벼운 퍼짐성 에스터.", "—", "일부 여드름유발 보고 — 지성은 소량"),
    # 항산화 ----------------------------------------------------------------
    ("항산화", "토코페롤 (Vitamin E)", "0.1~1%",
     "오일·제형 산패 방지(필수). 피부 항산화.", "자체가 천연", "과량 시 끈적·알러지 드묾"),
    ("항산화", "로즈마리추출물(올레오레진)/BHT", "0.05~0.5%",
     "오일 산패 지연. 무수·오일 제형 보존기간↑.", "로즈마리(천연)", "BHT는 합성 — 천연시 로즈마리"),
]

# ── B) 비율·제형 민감도 사다리 ────────────────────────────────────────────
# (시스템, [(구간, 결과/제형·상품성, 팁), ...])
LADDERS = [
    ("유화제 %(O/W)", [
        ("0%", "오일+물 → 섞여도 곧 분리. 2상 미스트/오일만 가능", "흔들어 쓰는 2층 미스트는 의도적 무유화"),
        ("2~3%", "묽은 로션(에멀전). 가벼운 발림", "공유화(세틸알콜 1~2%)로 점도·안정 보완"),
        ("4~5%", "로션~크림 경계. 데일리 보습", "히어로 montanov-68도 이 구간"),
        ("6~8%", "진한 크림. 리치·오클루시브", "과다 시 비누화·답답함·백탁 주의"),
    ]),
    ("왁스/경화 %(무수 밤)", [
        ("5~10%", "소프트밤·버터. 말랑 떠먹는 제형", "손 체온에 녹는 질감"),
        ("10~18%", "밤·연고. 적당히 단단", "밀랍/칸데릴라 비율로 미세조정"),
        ("18~25%", "스틱·립밤. 밀어 올려 바름", "칸데릴라는 밀랍보다 단단(식물성)"),
        ("25%+", "데오/고형 스틱. 매우 단단", "과다 시 끌림·백탁·발림 저하"),
    ]),
    ("계면활성제 %(세정)", [
        ("5~10%", "순한 저자극 워시·클렌징워터", "아미노산/글루코사이드계로 순하게"),
        ("10~20%", "일반 폼·젤 클렌저·샴푸", "약산성(5~6) 맞추면 부담↓"),
        ("20~40%", "진한 거품·고세정 폼/샴푸", "과하면 건조·당김 — 보습·베타인 보완"),
    ]),
    ("점증제 %(잔탄/카보머)", [
        ("잔탄 0.2~0.5%", "가벼운 점도·유화 안정", "글리세린 선분산 후 투입(덩어리 방지)"),
        ("잔탄 0.5~1%", "젤~세럼 바디감", ">0.7% 실뭉침·늘어짐 주의"),
        ("카보머 0.1~0.3%", "맑고 가벼운 젤", "중화 필요(TEA/NaOH)"),
        ("카보머 0.3~0.5%", "탄탄한 투명 젤", "전해질·저pH·알콜에 무너짐"),
    ]),
    ("습윤제 %(글리세린)", [
        ("2~5%", "적정 보습. 끈적임 적음", "대부분 제형의 표준 구간"),
        ("5~8%", "고보습·촉촉막", "환경 건조 시 역삼투(수분 빼앗김) 주의"),
        ("8%+", "끈적임·무거움", "스프레이보다 펌핑 용기 권장"),
    ]),
    ("유상(오일) %", [
        ("0%", "토너·미스트(수상 85~95%)", "가용화제로 소량 오일/EO만"),
        ("5~15%", "로션. 가벼운 보습막", "유화제 3~5% 동반"),
        ("15~30%", "크림. 리치 보습", "유화·점도 설계 중요"),
        ("무수(100%)", "오일·밤. 방부 불필요", "물 유입 금지·항산화 필수"),
    ]),
    ("향/에센셜오일 %", [
        ("0.1~0.5%", "은은한 향(leave-on 안전권)", "가용화 없이 물에 넣으면 분리"),
        ("0.5~1%", "뚜렷한 향(바디 등)", "피부 자극·광독성(감귤계) 주의"),
        (">1%", "권장 상한 초과", "자극·규제 위험 — 지양"),
    ]),
]


def _style_header(ws, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(1, i, t)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor="E8EEF7")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w


def build(out: Path = Path("원료_지식_확장.xlsx")) -> Path:
    thin = Side(style="thin", color="D8DCE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    note = ("⚠ 업계 통용 배합지식(예시). 전형%·고시·안전성·방부 적정성은 집행 전 "
            "식약처 고시·공급사 TDS·방부력(챌린지) 테스트로 반드시 검증.")

    wb = openpyxl.Workbook()
    # 시트1: 기능성 원료 팔레트
    ws = wb.active
    ws.title = "기능성원료"
    ws.merge_cells("A1:F1")
    n = ws.cell(1, 1, note)
    n.font = Font(bold=True, size=9, color="7A4A2A")
    n.fill = PatternFill("solid", fgColor="FDECE0")
    n.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 26
    titles = ["역할", "원료 (INCI)", "전형%", "기능 · 왜 쓰나", "천연 대안", "주의"]
    widths = [10, 34, 9, 46, 20, 40]
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(2, i, t)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor="E8EEF7")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 3
    for role, name, pct, func, nat, caution in PALETTE:
        vals = [role, name, pct, func, nat, caution]
        for ci, v in enumerate(vals, 1):
            cc = ws.cell(r, ci, v)
            cc.alignment = wrap
            cc.border = border
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=ROLE_COLORS.get(role, "F0F0F0"))
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="top")
        longest = max(len(str(func)), len(str(caution)) * 46 // 40)
        ws.row_dimensions[r].height = min(120, max(28, -(-longest // 23) * 15))
        r += 1
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:F{r - 1}"

    # 시트2: 비율·제형 사다리
    ws2 = wb.create_sheet("비율_제형_사다리")
    ws2.merge_cells("A1:D1")
    n2 = ws2.cell(1, 1, "비율이 제형·상품성을 가른다 — 같은 원료군도 %에 따라 결과가 바뀜. "
                        + note)
    n2.font = Font(bold=True, size=9, color="7A4A2A")
    n2.fill = PatternFill("solid", fgColor="FDECE0")
    n2.alignment = Alignment(wrap_text=True, vertical="center")
    ws2.row_dimensions[1].height = 30
    _style_header2 = ["시스템", "구간(%)", "결과 · 제형/상품성", "팁"]
    widths2 = [20, 16, 44, 40]
    for i, (t, w) in enumerate(zip(_style_header2, widths2), 1):
        c = ws2.cell(2, i, t)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor="E8EEF7")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws2.column_dimensions[get_column_letter(i)].width = w
    r = 3
    band = PatternFill("solid", fgColor="F7FAFF")
    for si, (system, steps) in enumerate(LADDERS):
        start = r
        for seg, result, tip in steps:
            ws2.cell(r, 1, system if r == start else "")
            for ci, v in enumerate(["", seg, result, tip], 1):
                if ci == 1:
                    continue
                cc = ws2.cell(r, ci, v)
                cc.alignment = wrap
                cc.border = border
            ws2.cell(r, 1).border = border
            if si % 2 == 1:
                for ci in range(1, 5):
                    ws2.cell(r, ci).fill = band
            ws2.row_dimensions[r].height = max(24, -(-len(str(result)) // 22) * 15)
            r += 1
        ws2.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws2.cell(start, 1).font = Font(bold=True, size=10)
        ws2.cell(start, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws2.freeze_panes = "A3"

    wb.save(out)
    print(f"원료 지식 확장 → {out.name}: 기능성원료 {len(PALETTE)}종 · "
          f"사다리 {len(LADDERS)}시스템")
    return out


if __name__ == "__main__":
    build()
