"""인쇄·제본용 책 PDF 생성 — 원료_마스터.xlsx → 1권, 제품개발_마스터.xlsx → 2권 (B5·스프링 제본).

엑셀 내용을 그대로 읽어(수정 없이) 책 레이아웃으로 다시 앉힌다.
  - 시트 = 장(chapter), 시트 묶음 = 부(part). 표지·목차·부 표지·쪽번호·머리말 자동.
  - 1행=장 부제, 2행=도입 문단. 남색 머리글 행=표, 단일 셀 행=제목/문단(서식으로 수준 판별).
  - 원료 카드(33열)는 원료 1개=반 쪽 카드로, 원료 도감은 2단 압축 항목으로 재배치.
  - 수식·입력 시트(처방설계·진도판·개발일지·입출고)는 손으로 채우는 기입 양식으로.
  - 장 표지·앞표지 일러스트는 인쇄본/assets/illust/<키>.png 가 있으면 넣고, 없으면 글자만.

실행:  uv run --with weasyprint python scripts/print_book.py [--vol 1|2] [--html-only]
산출:  인쇄본/1권_원료마스터.pdf, 인쇄본/2권_제품개발마스터.pdf (+ 같은 이름 .html)
"""

from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import book_diagrams

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "인쇄본"
ASSETS = OUT / "assets"
ILLUST = ASSETS / "illust"
EMBED = ROOT / "제품_공정이미지" / "_embed"

DARK_FILLS = {"FF1F3A5F", "FF33475B", "FF2F5597"}
H2_FILLS = {"FFDDEBF7"}
H3_FILLS = {"FFDCE6F1"}

# ── 책 구성 ────────────────────────────────────────────────────────────────
VOLUMES = {
    1: dict(
        file="원료_마스터.xlsx",
        out="1권_원료마스터",
        vol="제1권",
        title="원료 마스터",
        subtitle="원료 하나를 끝까지 파서 원하는 제형을 만든다",
        parts=[
            ("들어가며", None, ["읽어보기"]),
            ("Ⅰ", "원료 카드 247", ["원료 카드"]),
            ("Ⅱ", "비교와 설계", ["비교 매트릭스", "제형 설계", "감촉 레버"]),
            ("Ⅲ", "넓혀 가기", ["블렌드 해부", "다음 공부"]),
            ("부록", None, ["진도판", "@원료 색인"]),
        ],
    ),
    2: dict(
        file="제품개발_마스터.xlsx",
        out="2권_제품개발마스터",
        vol="제2권",
        title="제품개발 마스터",
        subtitle="집 개발에서 OEM 출시까지 — 레시피·기록·규제·브랜드·시험",
        parts=[
            ("들어가며", None, ["읽어보기"]),
            ("Ⅰ", "만들기", ["레시피", "역설계 심화", "처방설계", "원료 도감"]),
            ("Ⅱ", "기록하고 고치기", ["개발일지", "트러블슈팅", "원료", "입출고", "도구"]),
            ("Ⅲ", "규제와 양산", ["사용제한원료", "유통안전·OEM검수", "OEM 개발의뢰서", "개발→창업 워크플로우"]),
            ("Ⅳ", "브랜드와 마케팅", ["브랜드 기획서", "수요검증", "업로드 루틴", "표시·광고 컴플라이언스"]),
            ("Ⅴ", "시험 대비", ["교재 이론 목차", "시험 핵심개념", "시험대비 원료맵", "시험→사업 활용", "실습↔시험 개념"]),
            ("부록", None, ["처방"]),
        ],
    ),
}

# 장 도입부 뒤에 넣는 개념 도식 (book_diagrams.py)
DIAGRAMS = {
    (1, "읽어보기"): ["daily_loop"],
    (1, "원료 카드"): ["card_anatomy"],
    (1, "제형 설계"): ["emulsion"],
    (1, "감촉 레버"): ["texture_levers"],
    (2, "읽어보기"): ["dev_loop", "stock_flow"],
    (2, "레시피"): ["three_phase"],
    (2, "역설계 심화"): ["ratio_ladder"],
    (2, "개발→창업 워크플로우"): ["ten_steps"],
    (2, "수요검증"): ["demand_stairs"],
}

RECIPE_IMAGES = ["01_토너", "02_앰플", "03_젤", "04_딥클렌징오일", "05_워시오프", "06_로션", "07_크림", "08_미니멀크림"]


# ── 셀 읽기 ────────────────────────────────────────────────────────────────
@dataclass
class Cell:
    col: int
    text: str
    bold: bool
    fill: str | None
    size: float | None
    white: bool = False


def _rgb(color) -> str | None:
    try:
        v = color.rgb
        return v if isinstance(v, str) else None
    except Exception:
        return None


def read_rows(ws, start=3):
    """(행번호, [Cell...]) — 값 있는 셀만. 수식 결과가 없는 셀(=로 시작)은 빈칸 취급."""
    for r in range(start, ws.max_row + 1):
        cells = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if v is None or (isinstance(v, str) and (not v.strip() or v.startswith("="))):
                continue
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            if hasattr(v, "strftime"):
                v = v.strftime("%Y-%m-%d")
            fill = _rgb(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else None
            white = _rgb(cell.font.color) == "FFFFFFFF" if cell.font.color else False
            cells.append(Cell(c - 1, str(v).strip(), bool(cell.font.b), fill, cell.font.sz, white))
        yield r, cells


def esc(s: str) -> str:
    return html.escape(s).replace("\n", "<br>")


def col_width(ws, idx: int) -> float:
    d = ws.column_dimensions.get(get_column_letter(idx + 1))
    return float(d.width) if d and d.width else 9.0


# ── 블록 모델 ──────────────────────────────────────────────────────────────
@dataclass
class Table:
    cols: list[int]
    heads: list[str]
    widths: list[float]
    rows: list[tuple[list[str], bool, str | None]] = field(default_factory=list)  # (cells, bold첫칸, fill)

    def finish(self, keep_empty=False):
        """데이터가 하나도 없는 열(수식 결과 없음)은 뺀다. 폭 미지정 열은 내용 길이로 보정."""
        if not self.rows or keep_empty:
            return self
        for i in range(len(self.cols)):
            lens = [len(r[0][i]) for r in self.rows if r[0][i]]
            avg = sum(lens) / len(lens) if lens else 0
            self.widths[i] = max(self.widths[i], min(36.0, avg * 0.8))
        keep = [i for i in range(len(self.cols)) if any(r[0][i] for r in self.rows)]
        keep = keep or [0]
        self.cols = [self.cols[i] for i in keep]
        self.heads = [self.heads[i] for i in keep]
        self.widths = [self.widths[i] for i in keep]
        self.rows = [([r[0][i] for i in keep], r[1], r[2]) for r in self.rows]
        return self


def table_html(t: Table, extra_class="", blank_rows=0) -> str:
    total = sum(t.widths) or 1
    # 표 성격: 열 폭 합(엑셀 문자 단위)이 크면 글자 작게, 아주 크면 레코드형
    cls = "t-s" if total > 120 else ("t-m" if total > 85 else "")
    if total > 175 and len(t.cols) >= 6 and t.rows and not blank_rows:
        return records_html(t)
    out = [f'<table class="dt {cls} {extra_class}"><colgroup>']
    out += [f'<col style="width:{w / total * 100:.1f}%">' for w in t.widths]
    out.append("</colgroup><thead><tr>")
    out += [f"<th>{esc(h)}</th>" for h in t.heads]
    out.append("</tr></thead><tbody>")
    for cells, bold, fill in t.rows:
        rc = ' class="tot"' if (cells[0] in ("합계", "계")) else ""
        out.append(f"<tr{rc}>")
        for i, v in enumerate(cells):
            c = ' class="b"' if (i == 0 and bold) else ""
            out.append(f"<td{c}>{esc(v)}</td>")
        out.append("</tr>")
    for _ in range(blank_rows):
        out.append("<tr class='blank'>" + "<td></td>" * len(t.cols) + "</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def records_html(t: Table) -> str:
    out = ['<div class="records">']
    for cells, _, _ in t.rows:
        out.append('<div class="rec">')
        out.append(f'<div class="rec-h">{esc(cells[0])}</div><dl>')
        for h, v in zip(t.heads[1:], cells[1:]):
            if v and v not in ("—", "-"):
                out.append(f"<dt>{esc(h)}</dt><dd>{esc(v)}</dd>")
        out.append("</dl></div>")
    out.append("</div>")
    return "".join(out)


def kv_html(rows) -> str:
    out = ['<table class="kv">']
    for cells in rows:
        a = cells[0]
        rest = cells[1:]
        if a.text in ("·", "•"):
            out.append(f'<tr><td colspan="2" class="bul">{esc(" ".join(c.text for c in rest))}</td></tr>')
            continue
        bc = " b" if a.bold else ""
        out.append(f'<tr><th class="{bc}">{esc(a.text)}</th><td>{"<br>".join(esc(c.text) for c in rest)}</td></tr>')
    out.append("</table>")
    return "".join(out)


class Toc:
    def __init__(self):
        self.items: list[tuple[int, str, str]] = []  # (level, text, anchor)
        self.n = 0

    def anchor(self, level, text) -> str:
        self.n += 1
        a = f"a{self.n}"
        self.items.append((level, text, a))
        return a


# ── 일반 시트 렌더러 ───────────────────────────────────────────────────────
def render_generic(ws, start=3, hooks=None, toc: Toc | None = None, toc_h2=True) -> str:
    hooks = hooks or {}
    parts: list[str] = []
    table: Table | None = None
    last_header = None
    kv: list = []

    def flush():
        nonlocal table, kv
        if table is not None:
            parts.append(table_html(table.finish()))
            table = None
        if kv:
            parts.append(kv_html(kv))
            kv = []

    for r, cells in read_rows(ws, start):
        if not cells:
            flush()
            continue
        if "row" in hooks:
            h = hooks["row"](r, cells)
            if h is not None:
                flush()
                parts.append(h)
                continue
        first = cells[0]
        is_header = len(cells) >= 2 and all(c.bold and c.fill and (c.fill in DARK_FILLS or c.white) for c in cells)
        if is_header:
            flush()
            cols = [c.col for c in cells]
            table = Table(cols, [c.text for c in cells], [col_width(ws, c) for c in cols])
            last_header = (cols, [c.text for c in cells])
            continue
        # 제목 띠로 끊긴 뒤 같은 모양(3칸↑)의 행이 이어지면 직전 머리글로 표를 다시 연다
        if table is None and last_header and len(cells) >= 3 and cells[-1].col <= last_header[0][-1]:
            flush()
            cols, heads = last_header
            table = Table(cols, list(heads), [col_width(ws, c) for c in cols])
        if len(cells) == 1:
            flush()
            t = first.text
            if first.fill in DARK_FILLS:
                if "band" in hooks:
                    parts.append(hooks["band"](r, t))
                else:
                    a = toc.anchor(3, t) if (toc and toc_h2) else ""
                    parts.append(f'<h3 class="band" id="{a}">{esc(t)}</h3>')
            elif first.bold and (first.fill in H2_FILLS or (first.size or 11) >= 12):
                a = toc.anchor(3, t) if (toc and toc_h2) else ""
                parts.append(f'<h3 class="sec" id="{a}">{esc(t)}</h3>')
            elif first.bold and first.fill in H3_FILLS:
                parts.append(f'<h4 class="sub">{esc(t)}</h4>')
            elif first.bold and first.fill:
                parts.append(f'<h5 class="sub2">{esc(t)}</h5>')
            elif first.bold:
                parts.append(f'<p class="strong">{esc(t)}</p>')
            else:
                cls = "note" if (first.fill and first.fill not in ("FFFFFFFF",)) else ""
                if t.startswith(("⚠", "※")):
                    cls = "warn"
                parts.append(f'<p class="{cls}">{esc(t)}</p>')
            continue
        # 여러 칸
        if table is not None:
            row = [""] * len(table.cols)
            for c in cells:
                idx = max([i for i, hc in enumerate(table.cols) if hc <= c.col] or [0])
                row[idx] = (row[idx] + " " + c.text).strip() if row[idx] else c.text
            table.rows.append((row, first.bold, first.fill))
        else:
            kv.append(cells)
    flush()
    return "\n".join(parts)


# ── 장(시트)별 전용 렌더러 ─────────────────────────────────────────────────
def render_cards(ws, toc: Toc) -> str:
    """원료 카드 33열 → 원료 1개 = 반 쪽 카드. 역할별 절로 묶는다."""
    heads = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    groups = []  # (제목, [열 index])
    for m in sorted(ws.merged_cells.ranges, key=lambda m: m.min_col):
        if m.min_row == 3 and m.max_row == 3:
            groups.append((ws.cell(3, m.min_col).value, list(range(m.min_col - 1, m.max_col))))
    # 마지막 '✍ 내 기록' 그룹은 병합이 아닐 수 있다 → 30~32열
    last = max(i for _, cols in groups for i in cols)
    if last + 1 < len(heads):
        groups.append((ws.cell(3, last + 2).value or "✍ 내 기록", list(range(last + 1, len(heads)))))
    head_cols = {0, 1, 2, 3, 4}
    rec_group = groups[-1]
    body_groups = groups[:-1]

    by_role: dict[str, list] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] is None:
            continue
        by_role.setdefault(str(row[4]), []).append(row)

    out = []
    for role, rows in by_role.items():
        a = toc.anchor(3, f"{role} ({len(rows)})")
        out.append(f'<h3 class="sec role-h" id="{a}">{esc(role)} <span class="cnt">{len(rows)}종</span></h3>')
        for row in rows:
            v = ["" if x is None else str(x).strip() for x in row]
            own = v[3] == "✓"
            out.append(f'<div class="card" id="card-{v[0]}">')
            out.append(
                f'<div class="card-h"><span class="no">{esc(v[0])}</span>'
                f'<span class="nm">{esc(v[1])}</span>'
                f'<span class="inci">{esc(v[2])}</span>'
                f'<span class="chips"><span class="chip role">{esc(v[4])}</span>'
                f'{"<span class=\"chip own\">✓ 보유</span>" if own else ""}</span></div>'
            )
            blocks = []
            for gname, cols in body_groups:
                items = [(heads[i], v[i]) for i in cols if i not in head_cols and v[i] and v[i] not in ("—", "-", "–")]
                if not items:
                    continue
                h = f'<div class="grp"><div class="grp-t">{esc(str(gname))}</div><table class="gt">'
                h += "".join(f"<tr><th>{esc(str(k))}</th><td>{esc(val)}</td></tr>" for k, val in items)
                blocks.append((h + "</table></div>", sum(len(val) + 14 for _, val in items)))
            # 2칸 균형 배분(순서 유지): 앞에서부터 왼쪽에 쌓다가 절반을 넘으면 오른쪽
            total = sum(w for _, w in blocks)
            left, acc = [], 0
            for b, w in blocks:
                if acc + w / 2 <= total / 2 or not left:
                    left.append(b)
                    acc += w
                else:
                    break
            right = [b for b, _ in blocks[len(left):]]
            out.append(f'<table class="card-b"><tr><td>{"".join(left)}</td><td>{"".join(right)}</td></tr></table>')
            filled = {heads[i]: v[i] for i in rec_group[1] if v[i]}
            lv = filled.get("숙련도", "")
            boxes = "".join(
                f'<span class="ck{" on" if lv.startswith(str(k)) else ""}">{k} {t}</span>'
                for k, t in enumerate(["모름", "읽음", "만져봄", "처방에 씀"])
            )
            memo = esc(filled.get("내 관찰 메모", ""))
            hand = "☑" if filled.get("손테스트") else "☐"
            out.append(
                f'<div class="card-f"><span class="lbl">✍ 손테스트 {hand}</span>'
                f'<span class="lbl">숙련도</span>{boxes}<span class="memo">메모 {memo}</span></div>'
            )
            out.append("</div>")
    return "\n".join(out)


def render_card_index(wb) -> str:
    ws = wb["원료 카드"]
    rows = [(str(r[1]), str(r[2] or ""), r[0]) for r in ws.iter_rows(min_row=5, values_only=True) if r[1]]
    rows.sort(key=lambda x: x[0])
    out = ['<p class="lead">원료 카드 247종 가나다순. 오른쪽 숫자는 쪽.</p><div class="index">']
    initial = None
    for nm, inci, no in rows:
        ch = _initial(nm)
        if ch != initial:
            initial = ch
            out.append(f'<div class="ix-h">{esc(ch)}</div>')
        out.append(
            f'<a class="ix" href="#card-{no}"><span class="ix-n">{esc(nm)}</span>'
            f'<span class="ix-i">{esc(inci)}</span></a>'
        )
    out.append("</div>")
    return "".join(out)


_CHO = "ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ"


def _initial(s: str) -> str:
    c = s[0]
    if "가" <= c <= "힣":
        ch = _CHO[(ord(c) - 0xAC00) // 588]
        return {"ㄲ": "ㄱ", "ㄸ": "ㄷ", "ㅃ": "ㅂ", "ㅆ": "ㅅ", "ㅉ": "ㅈ"}.get(ch, ch)
    return "A–Z·기타" if not c.isdigit() else "0–9"


def _dg_name(s: str) -> tuple[str, str]:
    """'호호바오일\n(= 별칭 · …)' → ('호호바오일', '별칭 · …')"""
    name, _, alias = s.partition("\n")
    return name.strip(), alias.strip().removeprefix("(=").removesuffix(")").strip()


def _dg_short(name: str) -> str:
    # 머리말용 짧은 이름: 뒤쪽 '(INCI…)' 설명을 떼고 18자까지
    n = re.sub(r"\s*\([^)]*\)\s*$", "", name).strip() or name
    return n if len(n) <= 18 else n[:17] + "…"


def _dg_key(name: str) -> str:
    # 색인 정렬용: '(유기농)·(프리미엄)' 같은 앞 괄호는 떼고 본 이름으로
    return re.sub(r"^\([^)]*\)\s*", "", name).strip() or name


def render_dogam(ws, toc: Toc) -> str:
    """구분(기능성 F·허브누리 H) → 역할 순으로 다시 묶고 번호를 매긴다.
    시트에선 같은 역할이 여러 번 끊겨 나오지만, 책에선 역할당 한 덩어리.
    쪽 머리말에 그 쪽의 역할·번호 범위가 찍히고, 앞엔 역할 안내표, 뒤엔 가나다 색인."""
    heads = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    groups: dict[str, dict[str, list]] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        v = ["" if x is None else str(x).strip() for x in row]
        groups.setdefault(v[3], {}).setdefault(v[1] or "기타", []).append(v)

    # 번호 매기기: 기능성=F, 허브누리=H, 그 외 구분은 첫 글자
    numbered = []  # (group, role, no, v)
    for g, roles in groups.items():
        p = {"기능성": "F", "허브누리": "H"}.get(g, g[:1])
        width = 3 if sum(len(x) for x in roles.values()) >= 100 else 2
        i = 0
        for role, items in roles.items():
            for v in items:
                i += 1
                numbered.append((g, role, f"{p}{i:0{width}d}", v))

    out = []
    # 역할 안내표
    guide = ['<table class="dg-guide"><thead><tr><th>구분</th><th>역할</th><th>종수</th><th>번호</th><th>쪽</th></tr></thead><tbody>']
    for g, roles in groups.items():
        for role in roles:
            nos = [no for gg, rr, no, _ in numbered if gg == g and rr == role]
            guide.append(
                f'<tr><td>{esc(g)}</td><td><b>{esc(role)}</b></td><td>{len(nos)}</td>'
                f'<td>{nos[0]}{" – " + nos[-1] if len(nos) > 1 else ""}</td>'
                f'<td><a href="#dg-{nos[0]}"></a></td></tr>'
            )
    guide.append("</tbody></table>")
    out.append(
        '<p class="lead">원료마다 번호가 붙어 있다(기능성 F, 허브누리 H). 쪽 위 머리말에 그 쪽의 역할과 번호 범위가 나온다. '
        "역할로 찾을 땐 아래 안내표를, 이름으로 찾을 땐 이 장 끝의 가나다 색인을 본다.</p>"
    )
    out.append("".join(guide))
    out.append('<div class="dogam-body">')  # 여기부터 도감 전용 머리말(page: dogam)

    cur_group = cur_role = None
    for g, role, no, v in numbered:
        if g != cur_group:
            if cur_group is not None:
                out.append("</div>")
            cur_group, cur_role = g, None
            a = toc.anchor(3, f"{g} 원료")
            out.append(f'<h3 class="sec dg-sec" id="{a}">{esc(g)} 원료</h3><div class="dogam">')
        if role != cur_role:
            cur_role = role
            nos = [n for gg, rr, n, _ in numbered if gg == g and rr == role]
            rng = f"{nos[0]} – {nos[-1]}" if len(nos) > 1 else nos[0]
            out.append(
                f'<div class="dg-role"><span class="rn">{esc(g)} · {esc(role)}</span>'
                f'<span class="rc">{len(nos)}종 · {rng}</span></div>'
            )
        name, alias = _dg_name(v[0])
        chips = []
        if v[2]:
            chips.append(f"전형 {esc(v[2])}")
        if v[4]:
            chips.append(f"등장 {esc(v[4])}회")
        if v[10]:
            chips.append('<b class="own">' + esc(v[10]) + "</b>")
        out.append(
            f'<div class="dg" id="dg-{no}"><div class="dg-n" data-k="{no} {esc(_dg_short(name))}"><span class="no">{no}</span> {esc(name)}</div>'
            + (f'<div class="dg-a">= {esc(alias)}</div>' if alias else "")
        )
        if chips:
            out.append(f'<div class="dg-c">{" · ".join(chips)}</div>')
        if v[6]:
            out.append(f'<div class="dg-f">{esc(v[6])}</div>')
        for i, lab in ((5, "사용%"), (7, None), (8, "주의"), (9, "대체")):
            if v[i] and v[i] not in ("—", "-"):
                lab = lab or heads[i].split("/")[0].strip()
                out.append(f'<div class="dg-x"><i>{esc(lab)}</i> {esc(v[i])}</div>')
        out.append("</div>")
    if cur_group is not None:
        out.append("</div>")

    # 가나다 색인
    a = toc.anchor(3, "원료 도감 가나다 색인")
    out.append(f'<h3 class="sec dg-sec" id="{a}">가나다 색인 — 이름 → 번호 · 쪽</h3><div class="dg-index">')
    entries = sorted(((_dg_key(_dg_name(v[0])[0]), _dg_name(v[0])[0], no) for _, _, no, v in numbered),
                     key=lambda x: (x[0].lower(), x[2]))
    initial = None
    for key, name, no in entries:
        ch = _initial(key)
        if ch != initial:
            initial = ch
            out.append(f'<div class="ix-h">{esc(ch)}</div>')
        out.append(f'<a class="dgx" href="#dg-{no}" data-k="{esc(_dg_short(name))}"><span class="no">{no}</span> {esc(re.sub(r"\s*\([^)]*\)\s*$", "", name) or name)}</a>')
    out.append("</div></div>")
    return "\n".join(out)


def render_devlog(ws, toc: Toc) -> str:
    heads = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    heads = [h for h in heads if h]
    entries = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    out = [f'<h3 class="sec" id="{toc.anchor(3, "지금까지의 기록")}">지금까지의 기록</h3>', '<div class="logs">']
    for e in entries:
        v = ["" if x is None else (x.strftime("%Y-%m-%d") if hasattr(x, "strftime") else str(x)) for x in e]
        out.append('<div class="logcard">' + _logcard(heads, v) + "</div>")
    out.append("</div>")
    out.append(f'<h3 class="sec pb" id="{toc.anchor(3, "기록지 (빈 양식)")}">기록지 (빈 양식)</h3>')
    out.append('<p class="lead">한 배치 = 한 칸. 한 번에 한 가지만 바꾼다. 필요하면 이 쪽을 복사해서 쓴다.</p>')
    out.append('<div class="logs">')
    for _ in range(16):
        out.append('<div class="logcard empty">' + _logcard(heads, [""] * len(heads)) + "</div>")
    out.append("</div>")
    return "\n".join(out)


def _logcard(heads, v):
    m = dict(zip(heads, v))
    top = ["날짜", "제품", "배치량g", "기준(base)"]
    nums = ["유화ok", "분리mm", "pH", "판정"]
    lines = ["바꾼것", "사용감", "이슈", "다음 아이디어"]
    cell = lambda h: f'<td><i>{h}</i><span>{esc(m.get(h, ""))}</span></td>'
    s = '<table class="lc"><tr>' + "".join(cell(h) for h in top) + "</tr><tr>" + "".join(cell(h) for h in nums) + "</tr></table>"
    s += '<table class="lc-l">' + "".join(f'<tr><th>{h}</th><td>{esc(m.get(h, ""))}</td></tr>' for h in lines) + "</table>"
    return s


def render_ledger(ws, toc: Toc) -> str:
    heads = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    cols = [i for i, h in enumerate(heads) if h]
    t = Table(cols, [heads[i] for i in cols], [col_width(ws, i) for i in cols])
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0]:
            t.rows.append(([str(r[i] or "") for i in cols], False, None))
    return table_html(t, "form", blank_rows=max(0, 52 - len(t.rows)))


def render_progress(ws, toc: Toc) -> str:
    """진도판 — 수식 집계표를 손으로 세는 기입표로."""
    out = ['<h4 class="sub">전체 진도</h4><table class="dt form"><thead><tr><th>항목</th><th style="width:22%">값</th></tr></thead><tbody>']
    for r in range(5, 10):
        out.append(f"<tr><td>{esc(str(ws.cell(r, 1).value))}</td><td></td></tr>")
    out.append("</tbody></table>")
    out.append('<h4 class="sub">숙련도별 원료 수</h4><table class="dt form"><thead><tr><th>숙련도</th><th>원료 수</th><th>비율</th></tr></thead><tbody>')
    for r in range(5, 10):
        out.append(f"<tr><td>{esc(str(ws.cell(r, 4).value))}</td><td></td><td></td></tr>")
    out.append("</tbody></table>")
    heads = [ws.cell(12, c).value for c in (1, 2, 4, 5, 6)]
    out.append('<h4 class="sub">역할별 진도</h4><table class="dt form"><thead><tr>' + "".join(f"<th>{esc(str(h))}</th>" for h in heads) + "</tr></thead><tbody>")
    for r in range(13, ws.max_row + 1):
        if ws.cell(r, 1).value:
            out.append(f"<tr><td>{esc(str(ws.cell(r, 1).value))}</td>" + "<td></td>" * 4 + "</tr>")
    out.append("</tbody></table>")
    return "\n".join(out)


def render_designer(ws, toc: Toc) -> str:
    """처방설계 캔버스 → 빈 설계지 2장 + 참고 골격 15종."""
    roles = [ws.cell(r, 12).value for r in range(8, 20)]
    skel = [ws.cell(r, 12).value for r in range(22, 31)]
    checks = [
        "합계 100%",
        "물+오일인데 유화제 없음 → 분리",
        "물 있는데 방부제 없음 → 변질",
        "세이프 X 원료 → 히어로(말라세지아-세이프) 계열 불가",
        "재고 부족 → 배치량↓ 또는 구매",
        "⛔ 법적 한도 초과",
        "향료 포함 — 다도기는 무향 원칙",
    ]
    heads = ["원료명", "역할", "투입%", "그램", "상", "세이프", "현재고g", "법적 한도", "메모"]
    widths = [24, 13, 10, 9, 7, 10, 10, 14, 14]
    sheet = []
    sheet.append('<div class="ds-top"><span>① 제형 <u></u></span><span>② 배치량 <u class="s"></u> g</span><span>제품명·버전 <u></u></span><span>날짜 <u class="s"></u></span></div>')
    sheet.append('<table class="designer"><tr><td class="ds-main">')
    sheet.append('<table class="dt form ds"><colgroup>' + "".join(f'<col style="width:{w / sum(widths) * 100:.1f}%">' for w in widths) + "</colgroup>")
    sheet.append("<thead><tr>" + "".join(f"<th>{h}</th>" for h in heads) + "</tr></thead><tbody>")
    for _ in range(18):
        sheet.append("<tr class='blank'>" + "<td></td>" * len(heads) + "</tr>")
    sheet.append("<tr class='tot'><td>합계</td><td></td><td></td><td></td>" + "<td></td>" * 5 + "</tr></tbody></table>")
    sheet.append('</td><td class="ds-side">')
    sheet.append('<table class="dt form mini"><thead><tr><th>역할별 합계</th><th>%</th></tr></thead><tbody>' + "".join(f"<tr><td>{esc(str(x))}</td><td></td></tr>" for x in roles) + "</tbody></table>")
    sheet.append('<table class="dt form mini"><thead><tr><th>골격 비교</th><th>내 %</th></tr></thead><tbody>' + "".join(f"<tr><td>{esc(str(x))}</td><td></td></tr>" for x in skel) + "</tbody></table>")
    sheet.append('<div class="checks"><b>점검</b>' + "".join(f"<div>☐ {esc(c)}</div>" for c in checks) + "</div>")
    sheet.append("</td></tr></table>")
    one = "\n".join(sheet)
    warn = ws.cell(29, 1).value or ""
    out = [one, f'<p class="warn">{esc(warn)}</p>', '<div class="pb"></div>', one]
    out.append(f'<h3 class="sec pb" id="{toc.anchor(3, "참고 골격 — 제형 15종")}">{esc(str(ws.cell(41, 1).value))}</h3>')
    rows = []
    for r in range(43, ws.max_row + 1, 2):
        rows.append(([str(ws.cell(r, 1).value or ""), str(ws.cell(r, 2).value or ""), str(ws.cell(r + 1, 2).value or "")], True, None))
    t = Table([0, 1, 2], ["제형", "목표 골격(전형 %)", "배합 주의"], [14, 40, 46], rows)
    out.append(table_html(t))
    return "\n".join(out)


def recipe_hooks(toc: Toc):
    state = {"n": -1}

    def band(r, t):
        if t.startswith("["):
            state["n"] += 1
            a = toc.anchor(3, t.split("]")[-1].split("·")[0].strip() if "]" in t else t)
            lab, _, rest = t.partition("]")
            return f'<h3 class="recipe-h" id="{a}"><span class="lv">{esc(lab.strip("["))}</span>{esc(rest.strip())}</h3>'
        return f'<h3 class="band">{esc(t)}</h3>'

    def row(r, cells):
        first = cells[0]
        if first.text.startswith("[") and first.fill in DARK_FILLS:
            return band(r, first.text)  # '▶ 실험 공정 이미지' 칸은 버린다
        if first.text.startswith("[") and first.bold and first.fill:  # 확장 예고 밴드
            return f'<h3 class="recipe-h ext">{esc(first.text)}</h3>'
        if len(cells) == 1 and first.text == "만드는 법" and 0 <= state["n"] < len(RECIPE_IMAGES):
            img = EMBED / f"{RECIPE_IMAGES[state['n']]}.jpg"
            fig = ""
            if img.exists():
                fig = f'<figure class="proc"><img src="{img.as_uri()}"><figcaption>실험 공정 — {esc(RECIPE_IMAGES[state["n"]][3:])}</figcaption></figure>'
            return fig + '<h4 class="sub">만드는 법</h4>'
        return None

    return {"band": band, "row": row}


def troubleshoot_tail(ws) -> str:
    img = EMBED / "불량판정_레퍼런스.jpg"
    if img.exists():
        return f'<figure class="wide"><img src="{img.as_uri()}"><figcaption>불량 판정 레퍼런스</figcaption></figure>'
    return ""


# ── 조립 ───────────────────────────────────────────────────────────────────
def chapter_html(vol: int, wb, name: str, toc: Toc) -> str:
    if name.startswith("@"):
        title = name[1:]
        a = toc.anchor(2, title)
        return f'<section class="chapter"><h2 class="ch" id="{a}">{esc(title)}</h2>{render_card_index(wb)}</section>'
    ws = wb[name]
    a = toc.anchor(2, name)
    t1 = str(ws.cell(1, 1).value or "")
    sub = t1.split("—", 1)[1].strip() if "—" in t1 else (t1 if t1 != name else "")
    lead = str(ws.cell(2, 1).value or "") if not ws.cell(2, 2).value else ""
    head = f'<h2 class="ch" id="{a}">{esc(name)}</h2>'
    if sub and sub != name:
        head += f'<p class="ch-sub">{esc(sub)}</p>'
    start = 3
    if lead and not ws.cell(2, 1).font.b:
        head += f'<p class="lead">{_lead(lead)}</p>'
    else:
        start = 2
    for d in DIAGRAMS.get((vol, name), []):
        head += book_diagrams.render(d, wb)
    special = {
        "원료 카드": lambda: render_cards(ws, toc),
        "원료 도감": lambda: render_dogam(ws, toc),
        "개발일지": lambda: render_devlog(ws, toc),
        "입출고": lambda: render_ledger(ws, toc),
        "진도판": lambda: render_progress(ws, toc),
        "처방설계": lambda: render_designer(ws, toc),
    }
    if name in special:
        body = special[name]()
    elif name == "레시피":
        body = render_generic(ws, start, recipe_hooks(toc), toc)
    else:
        many = sum(1 for _, cells in read_rows(ws, start) if len(cells) == 1 and cells[0].bold and ((cells[0].size or 11) >= 12 or cells[0].fill in DARK_FILLS))
        body = render_generic(ws, start, toc=toc, toc_h2=(many <= 18 and name != "역설계 심화") or name == "역설계 심화")
        if name == "트러블슈팅":
            body += troubleshoot_tail(ws)
    cls = "chapter" + (" cards" if name == "원료 카드" else "") + (" dogam-ch" if name == "원료 도감" else "")
    return f'<section class="{cls}" data-ch="{esc(name)}">{head}{body}</section>'


def _lead(s: str) -> str:
    # '┃' 로 이어 붙인 컬럼 설명은 줄을 나눠 읽기 쉽게
    return "<br>".join(esc(p.strip()) for p in s.split("┃") if p.strip())


def part_html(num, title, chapters, key, toc: Toc) -> str:
    label = f"{num}부" if num not in ("들어가며", "부록") else num
    a = toc.anchor(1, f"{num}. {title}" if title else num)
    img = ILLUST / f"{key}.png"
    art = f'<div class="part-art"><img src="{img.as_uri()}"></div>' if img.exists() else '<div class="part-art none"></div>'
    lis = "".join(f"<li>{esc(c.lstrip('@'))}</li>" for c in chapters)
    big = esc(title or num)
    sm = esc(label if title else "")
    return f'<section class="part" id="{a}">{art}<div class="part-txt"><div class="part-no">{sm}</div><h1 class="part-t">{big}</h1><ul>{lis}</ul></div></section>'


def toc_html(toc: Toc) -> str:
    out = ['<section class="toc"><h1 class="toc-t">차례</h1><div class="toc-l">']
    for level, text, a in toc.items:
        if level > 3:
            continue
        out.append(f'<a class="t{level}" href="#{a}"><span class="tx">{esc(text)}</span></a>')
    out.append("</div></section>")
    return "".join(out)


def cover_html(cfg, key) -> str:
    img = ILLUST / f"{key}.png"
    art = f'<img class="cv-art" src="{img.as_uri()}">' if img.exists() else ""
    return (
        f'<section class="cover">{art}<div class="cv-box">'
        f'<div class="cv-vol">{esc(cfg["vol"])}</div><h1 class="cv-t">{esc(cfg["title"])}</h1>'
        f'<div class="cv-s">{esc(cfg["subtitle"])}</div>'
        f'<div class="cv-b">다도기 홈랩 · 1인 브랜드 개발 노트</div></div></section>'
    )


def build(vol: int, html_only=False, only=None):
    cfg = VOLUMES[vol]
    wb = openpyxl.load_workbook(ROOT / cfg["file"], data_only=True)
    toc = Toc()
    body = []
    for i, (num, title, chapters) in enumerate(cfg["parts"]):
        key = f"vol{vol}_part{i}"
        if not only:
            body.append(part_html(num, title, chapters, key, toc))
        for ch in chapters:
            if not only or ch in only:
                body.append(chapter_html(vol, wb, ch, toc))
    doc = (
        "<!doctype html><html lang='ko'><head><meta charset='utf-8'>"
        f"<title>{esc(cfg['title'])}</title><style>{css(cfg)}</style></head><body>"
        + cover_html(cfg, f"vol{vol}_cover")
        + '<section class="blankpage"></section>'
        + toc_html(toc)
        + "\n".join(body)
        + "</body></html>"
    )
    OUT.mkdir(exist_ok=True)
    name = cfg["out"] + ("_미리보기" if only else "")
    hp = OUT / f"{name}.html"
    hp.write_text(doc, encoding="utf-8")
    print("html:", hp)
    if html_only:
        return
    from weasyprint import HTML

    pdf = OUT / f"{name}.pdf"
    HTML(string=doc, base_url=str(OUT)).write_pdf(pdf)
    print("pdf :", pdf)


def css(cfg) -> str:
    fonts = ASSETS / "fonts"
    # .ttf = scripts/otf2ttf.py 변환본. OTF(CID-CFF)를 쓰면 맥 미리보기에서 ↔·℃·괄호 등이 사라진다
    ff = "".join(
        f"@font-face{{font-family:P;src:url('{(fonts / f'Pretendard-{n}.ttf').as_uri()}');font-weight:{w}}}"
        for n, w in [("Light", 300), ("Regular", 400), ("Medium", 500), ("SemiBold", 600), ("Bold", 700), ("ExtraBold", 800)]
    )
    return ff + (ASSETS / "book.css").read_text(encoding="utf-8").replace("{{BOOK}}", cfg["title"])


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--vol", type=int, choices=[1, 2])
    ap.add_argument("--html-only", action="store_true")
    ap.add_argument("--only", nargs="*", help="이 시트(장)만 빌드 — 레이아웃 확인용 _미리보기.pdf")
    a = ap.parse_args()
    for v in [a.vol] if a.vol else [1, 2]:
        build(v, a.html_only, a.only)
